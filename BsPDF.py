#################实现自动先把ui 文件转为python 文件#################开始
import qt_ui_to_py
qt_ui_to_py.runMain()
#################实现自动先把ui 文件转为python 文件#################结束
from PyQt5.QtGui import *
from PyQt5.QtWidgets import QMainWindow, QApplication, QWidget, QFileDialog,QAbstractItemView,QMessageBox
import configparser
import ui_main_bs
# import resourcefile
import csv
import re
import fitz
import os
import sys
import datetime
import math
from shutil import copyfile
import win32com.client
from copy import deepcopy
from pathlib import Path
# import pandas as pd
from pandas import read_excel
from openpyxl import load_workbook

class mainwindow(QMainWindow):
    """
    主程序窗口
    """

    def __init__(self, window):
        super().__init__()
        self.ui = ui_main_bs.Ui_MainWindow()
        self.ui.setupUi(self)
        self.load_data()
        self.show()
        self.ui_event()

    def ui_event(self):
        """
        主窗口事件
        :return:
        """
        self.ui.lineEdit_BsKey.returnPressed.connect(self.getPDFitem)
        self.ui.comboBox_PDFitem.currentIndexChanged.connect(self.getPDFPath2)
        self.ui.open_pushButton1.clicked.connect(self.Open_File2)  # 按鈕打開PDF檔
        self.ui.open_pushButton_art.clicked.connect(self.Open_artFile)  # 按鈕打開art PDF檔
        self.ui.add_pushButton2.clicked.connect(self.add_File)
        self.ui.action_dele.triggered.connect(self.deleteRow)
        self.ui.action_importxls.triggered.connect(self.import_seq)

        # self.ui.pushButton_scan2.clicked.connect(self.scan2)
        # self.ui.action_output.triggered.connect(self.output_path)
        # self.ui.open_pushButton1.clicked.connect(self.Open_File)     #按鈕打開PDF檔
        # self.ui.open_tableView1.clicked.connect(self.pdf_info)  # 按鈕打開PDF檔
        self.ui.comboBox3.currentTextChanged.connect(self.check_bleeding)

        self.ui.process_pushButton3.clicked.connect(self.batch_combine)
        self.ui.lineEdit_r3c1.textChanged.connect(self.check_divmod)
        self.ui.lineEdit_r3c2.textChanged.connect(self.check_divmod)

        self.ui.action2.triggered.connect(self.revised_bleeding)



    def import_seq(self):
        """
        暫時不處理此情況
        :return:
        """
        global dirpath
        self.ui.progressBar.setValue(0)
        filelist, _ = QFileDialog.getOpenFileNames(
            self,  # 父窗口对象
            "选择要處理的實數編號Excel檔",  # 标题
            # r"./",  # 起始目录
            dirpath,
            "數據類型 (*.xls;*.xlsx)"  # 选择类型过滤项，过滤内容在括号中
        )


        if filelist!=[]:
            print("test:",filelist)
            xls = read_excel(filelist[0])
            seqlist_w = ()
            for i in xls.values:
                if math.isnan(i[1]):
                    seqlist_w = seqlist_w + (int(i[0]),)
                    # print(int(i[0]))
                else:
                    for r in range(int(i[0]), int(i[1] + 1)):
                        seqlist_w = seqlist_w + (r,)
                        # print(r)
            self.seqlist=seqlist_w
            self.group=False
            self.seqcode=""
            # print(filelist)
            self.BsKey=os.path.basename(filelist[0]).split('.')[0]

            print(seqlist_w)
        else:
            print("文件名空")

    def deleteRow(self):
        #選擇多少用
        # selectrows=self.ui.process_tableView2.selectionModel().selectedRows()
        # for row in sorted(selectrows):
        #     model_view2.removeRow(row.row())


        model_view2.removeRow(self.ui.process_tableView2.currentIndex().row())

        self.ui.process_tableView2.setModel(model_view2)
        self.ui.process_tableView2.resizeColumnsToContents()
        self.ui.process_tableView2.resizeRowsToContents()

    def getPDFitem(self):
        # print("test:")
        self.ui.progressBar.setValue(0)

        self.ui.comboBox_PDFitem.clear()

        print("key2:")
        # print("key2:", self.ui.lineEdit_BsKey.text())
        self.read_lotus_LaserRoom(self.ui.lineEdit_BsKey.text())
        self.ui.add_pushButton2.setStyleSheet("background-color: rgb(232, 232, 232)")
        self.ui.add_pushButton2.setEnabled(True)

    def getPDFPath(self):
        # self.ui.listWidget_path.clear()
        BsKey=self.ui.lineEdit_BsKey.text()
        # print("key:",BsKey)
        joblot = BsKey.split("#")[0]
        item = self.ui.comboBox_PDFitem.currentText()
        print("item:",item)
        if item != '':
            joblotitem=joblot+'-'+item
            # s = win32com.client.Dispatch('Notes.NotesSession')
            # db = s.GetDatabase("HMP03/IT/HMP", "PublicNSF\LaserMat.nsf")
            # view = db.GetView("searchPDFbyjoblotitem")
            dc = view_joblotitem.GetAllDocumentsByKey(joblotitem, True)
            print("dc.Count",dc.count,joblotitem)
            for i in range(1, dc.count + 1):
                doc = dc.GetNthDocument(i)
                path=doc.GetItemValue('path')[0]
                print("path2: ",path)
                if str(path.split(" ")[-1]).lower()[:4] == 'main':
                    LN_UserName=doc.GetItemValue('LN_UserName')[0]
                    print_type1=doc.GetItemValue('print_type')[0]   #編號次序
                    direction_type = doc.GetItemValue('direction_type')[0]  #PDF輸出方向
                    bleeding_type = doc.GetItemValue('bleeding_type')[0]    #出血位及顯示位置
                    PDF_rotate = doc.GetItemValue('PDF_rotate')[0]  # PDF角度
                    if PDF_rotate=='':
                        PDF_rotate='0'
                    else:
                        PDF_rotate=str(int(PDF_rotate))

                    check_Double_sided = doc.GetItemValue('check_Double_sided')[0]  #雙面打印

                    check_addblank = doc.GetItemValue('check_addblank')[0]      ##加隔紙
                    check_addblank2 = doc.GetItemValue('check_addblank2')[0]      #單面加稿件底頁
                    if check_addblank=='':
                        check_addblank='False'
                    if check_addblank2 == '':
                        check_addblank2 = 'False'
                    e1 = doc.GetItemValue('e1')[0]
                    e2 = str(float(doc.GetItemValue('e2')[0]))
                    e3 = str(float(doc.GetItemValue('e3')[0]))
                    e4 = doc.GetItemValue('e4')[0]
                    e5 = str(float(doc.GetItemValue('e5')[0]))
                    e6 = str(float(doc.GetItemValue('e6')[0]))
                    e7 = doc.GetItemValue('e7')[0]
                    e8 = doc.GetItemValue('e8')[0]

                    r1c5 = str(doc.GetItemValue('r1c5')[0])     #右邊增減
                    r1c6 = str(doc.GetItemValue('r1c6')[0])     #下邊增減
                    r2c3 = doc.GetItemValue('r2c3')[0]          #實際成品寬
                    r2c4 = doc.GetItemValue('r2c4')[0]          #實際成品高
                    r3c3 = str(doc.GetItemValue('r3c3')[0])     #水平位置
                    r3c4 = str(doc.GetItemValue('r3c4')[0])     #垂直位置
                    times_input = doc.GetItemValue('times_input')[0]
                    if times_input=='':
                        times_input='1'
                    else:
                        times_input=str(int(times_input))

                    width = doc.GetItemValue('width')[0]    #單個PDF檔寬度
                    height = doc.GetItemValue('height')[0]  ##單個PDF檔高度
                    path = doc.GetItemValue('path')[0]
                    print("path:",path)
                    # self.ui.listWidget_path.addItem(path)
                    filelist=[]
                    if LN_UserName=='CN=Ddata/O=HMP':
                        file = '//172.16.3.6/data/outgoing' + path[2:]
                    elif LN_UserName=='CN=DLaser/O=HMP':
                        # file = path
                        file='//172.16.3.6/data/outgoing' + path[2:]

                    filelist.append(file)
                    model = self.get_pdf_info(filelist)
                    self.ui.open_tableView1.setModel(model)
                    self.ui.open_tableView1.resizeColumnToContents(0)
                    self.ui.open_tableView1.resizeColumnToContents(1)
                    self.ui.open_tableView1.resizeRowsToContents()
            if not self.ui.open_tableView1.model():
                if self.ui.open_tableView1.model().rowCount() <= 1:
                    self.ui.comboBox1.setCurrentIndex(int(print_type1[0])-1)
                    self.ui.comboBox2.setCurrentIndex(int(direction_type[0])-1)
                    self.ui.comboBox3.setCurrentIndex(int(bleeding_type[0]))
                    self.ui.comboBox_rotate.setCurrentText(PDF_rotate)
                    self.ui.checkBox_blank.setChecked(eval(check_addblank))
                    self.ui.checkBox_blank2.setChecked(eval(check_addblank2))
                    self.ui.lineEdit_r1c1.setText(e1)
                    self.ui.lineEdit_r1c2.setText(e2)
                    self.ui.lineEdit_r1c3.setText(e3)
                    self.ui.lineEdit_r2c1.setText(e4)
                    self.ui.lineEdit_r2c2.setText(e5)
                    self.ui.lineEdit_r1c4.setText(e6)
                    self.ui.lineEdit_r3c1.setText(e7)
                    self.ui.lineEdit_r3c2.setText(e8)
                    self.ui.lineEdit_r1c5.setText(r1c5)
                    self.ui.lineEdit_r1c6.setText(r1c6)
                    self.ui.lineEdit_r2c3.setText(r2c3)
                    self.ui.lineEdit_r2c4.setText(r2c4)
                    self.ui.lineEdit_r3c3.setText(r3c3)
                    self.ui.lineEdit_r3c4.setText(r3c4)
                    self.ui.lineEdit_Repeat.setText(times_input)
                    self.ui.checkBox1.setChecked(eval(check_Double_sided))

                else:
                    QMessageBox.information(ui_mainwindow, "錯誤提示!", "不能有多個PDF檔!")

    def getPDFPath2(self):
        # self.ui.listWidget_path.clear()
        BsKey=self.ui.lineEdit_BsKey.text()
        # print("key:",BsKey)
        joblot = BsKey.split("#")[0]
        item = self.ui.comboBox_PDFitem.currentText()
        print("item:",item)
        if item != '':
            joblotitem=joblot+'-'+item
            # s = win32com.client.Dispatch('Notes.NotesSession')
            # db = s.GetDatabase("HMP03/IT/HMP", "PublicNSF\LaserMat.nsf")
            # view = db.GetView("searchPDFbyjoblotitem")
            dc = view_joblotitem.GetAllDocumentsByKey(joblotitem, True)
            print("dc.Count",dc.count,joblotitem)
            for i in range(1, dc.count + 1):
                doc = dc.GetNthDocument(i)
                path=doc.GetItemValue('path')[0]
                print("path2: ",path)
                if str(path.split(" ")[-1]).lower()[:4] == 'main':
                    LN_UserName=doc.GetItemValue('LN_UserName')[0]
                    print_type1=doc.GetItemValue('print_type')[0]   #編號次序
                    direction_type = doc.GetItemValue('direction_type')[0]  #PDF輸出方向
                    bleeding_type = doc.GetItemValue('bleeding_type')[0]    #出血位及顯示位置
                    PDF_rotate = doc.GetItemValue('PDF_rotate')[0]  # PDF角度
                    if PDF_rotate=='':
                        PDF_rotate='0'
                    else:
                        PDF_rotate=str(int(PDF_rotate))

                    check_Double_sided = doc.GetItemValue('check_Double_sided')[0]  #雙面打印

                    check_addblank = doc.GetItemValue('check_addblank')[0]      ##加隔紙
                    check_addblank2 = doc.GetItemValue('check_addblank2')[0]      #單面加稿件底頁
                    if check_addblank=='':
                        check_addblank='False'
                    if check_addblank2 == '':
                        check_addblank2 = 'False'
                    e1 = doc.GetItemValue('e1')[0]
                    e2 = str(float(doc.GetItemValue('e2')[0]))
                    e3 = str(float(doc.GetItemValue('e3')[0]))
                    e4 = doc.GetItemValue('e4')[0]
                    e5 = str(float(doc.GetItemValue('e5')[0]))
                    e6 = str(float(doc.GetItemValue('e6')[0]))
                    e7 = doc.GetItemValue('e7')[0]
                    e8 = doc.GetItemValue('e8')[0]

                    r1c5 = str(doc.GetItemValue('r1c5')[0])     #右邊增減
                    r1c6 = str(doc.GetItemValue('r1c6')[0])     #下邊增減
                    r2c3 = doc.GetItemValue('r2c3')[0]          #實際成品寬
                    r2c4 = doc.GetItemValue('r2c4')[0]          #實際成品高
                    r3c3 = str(doc.GetItemValue('r3c3')[0])     #水平位置
                    r3c4 = str(doc.GetItemValue('r3c4')[0])     #垂直位置
                    times_input = doc.GetItemValue('times_input')[0]
                    if times_input=='':
                        times_input='1'
                    else:
                        times_input=str(int(times_input))

                    width = doc.GetItemValue('width')[0]    #單個PDF檔寬度
                    height = doc.GetItemValue('height')[0]  ##單個PDF檔高度
                    path = doc.GetItemValue('path')[0]
                    print("path:",path)
                    # self.ui.listWidget_path.addItem(path)
                    #不用自動找filepath
                    # filelist=[]
                    # if LN_UserName=='CN=Ddata/O=HMP':
                    #     file = '//172.16.3.6/data/outgoinga' + path[2:]
                    # elif LN_UserName=='CN=DLaser/O=HMP':
                    #     # file = path
                    #     file='//172.16.3.6/data/outgoinga' + path[2:]
                    #
                    # filelist.append(file)
                    # 不用自動找filepath
                    # model = self.get_pdf_info(filelist)
                    # self.ui.open_tableView1.setModel(model)
                    # self.ui.open_tableView1.resizeColumnToContents(0)
                    # self.ui.open_tableView1.resizeColumnToContents(1)
                    # self.ui.open_tableView1.resizeRowsToContents()
            # if not self.ui.open_tableView1.model():
            #     if self.ui.open_tableView1.model().rowCount() <= 1:
                    self.ui.comboBox1.setCurrentIndex(int(print_type1[0])-1)
                    self.ui.comboBox2.setCurrentIndex(int(direction_type[0])-1)
                    self.ui.comboBox3.setCurrentIndex(int(bleeding_type[0]))
                    self.ui.comboBox_rotate.setCurrentText(PDF_rotate)
                    self.ui.checkBox_blank.setChecked(eval(check_addblank))
                    self.ui.checkBox_blank2.setChecked(eval(check_addblank2))
                    self.ui.lineEdit_r1c1.setText(e1)
                    self.ui.lineEdit_r1c2.setText(e2)
                    self.ui.lineEdit_r1c3.setText(e3)
                    self.ui.lineEdit_r2c1.setText(e4)
                    self.ui.lineEdit_r2c2.setText(e5)
                    self.ui.lineEdit_r1c4.setText(e6)
                    self.ui.lineEdit_r3c1.setText(e7)
                    self.ui.lineEdit_r3c2.setText(e8)
                    self.ui.lineEdit_r1c5.setText(r1c5)
                    self.ui.lineEdit_r1c6.setText(r1c6)
                    self.ui.lineEdit_r2c3.setText(r2c3)
                    self.ui.lineEdit_r2c4.setText(r2c4)
                    self.ui.lineEdit_r3c3.setText(r3c3)
                    self.ui.lineEdit_r3c4.setText(r3c4)
                    self.ui.lineEdit_Repeat.setText(times_input)
                    self.ui.checkBox1.setChecked(eval(check_Double_sided))

                else:
                    pass
                    # QMessageBox.information(ui_mainwindow, "錯誤提示!", "不能有多個PDF檔!")

    def read_lotus_LaserRoom_PDFItems(self,BsKey):
        key=BsKey.split("#")
        s = win32com.client.Dispatch('Notes.NotesSession')
        db = s.GetDatabase(ui_mainwindow.Lotus_server, "PublicNSF\\LaserMat.nsf")
        view = db.GetView("searchPDFbyjoblot")
        dc = view.GetAllDocumentsByKey(key[0], True)
        item =[]
        print("count_read:", dc.count)
        for i in range(1, dc.count+1):
            doc = dc.GetNthDocument(i)
            item.append(doc.GetItemValue('Item')[0])

        items=set(item)

        for k in items:
            print("items:", k)
            self.ui.comboBox_PDFitem.addItem(k)


    def read_lotus_LaserRoom(self,BsKey):
        key=BsKey.split("#")
        s = win32com.client.Dispatch('Notes.NotesSession')
        db = s.GetDatabase(self.Lotus_server, "PublicNSF\\LaserMat.nsf")
        db_bsseq = s.GetDatabase(self.Lotus_server, "PublicNSF\\QM.nsf")
        view = db.GetView("searchPDFbyjoblot")
        view_bsseq = db_bsseq.GetView("BsSeqBykey3")
        dc = view.GetAllDocumentsByKey(key[0], True)
        item =[]



        for i in range(1, dc.count+1):
            doc = dc.GetNthDocument(i)
            item.append(doc.GetItemValue('Item')[0])

        items=set(item)
        print(item)

        for k in items:
            print("items:", k)
            self.ui.comboBox_PDFitem.addItem(k)

        dc_bsseq = view_bsseq.GetAllDocumentsByKey(BsKey, True)
        doc_seqcode=dc_bsseq.GetNthDocument(1)
        self.seqcode = str(doc_seqcode.GetItemValue('id_code')[0]).upper().strip()

        print("self.seqcode:", self.seqcode)
        seqlist= ()
        seqlist2 = ()
        seqlist1 = []
        for j in range(1, dc_bsseq.count+1):
            doc_bsseq = dc_bsseq.GetNthDocument(j)
            seq = doc_bsseq.GetItemValue('num1')
            groupseq = doc_bsseq.GetItemValue('groupseq')
            print('Group: ',groupseq)
            # seqlist2=(seq,groupseq[0]+float('0.'+seq[0]))
            seqlist2 = (seq, groupseq[0] + float('0.' + str(10000000 + int(seq[0]))))
            # seqlist=seqlist+doc_bsseq.GetItemValue('num1')
            seqlist = seqlist + seq
            seqlist1.append(seqlist2)


        sort1 = sorted(seqlist1, key=lambda x: float(x[1]))
        sort2 = ()
        for s1 in sort1:
            sort2 = sort2 + (int(s1[0][0]),)
        print("seq:", seqlist)
        seqlist_w=()
        for sl1 in seqlist:
            seqlist_w=seqlist_w+(int(sl1),)

        if groupseq[0]==0:  #不是分組的.
            self.seqlist=sorted(seqlist_w)
            self.group=False
            self.BsKey = '' #匯入編號就有檔案名
        else:               #分組的, 用分組的排序
            self.seqlist = sort2
            self.group = True
            self.BsKey=''   #匯入編號就有檔案名
        print("seqsort:", self.seqlist)

    def load_data(self):
        # global outputfile
        self.ui.progressBar.setValue(0)
        self.ui.progressBar.setFormat("%p%")
        self.ui.add_pushButton2.setStyleSheet("background-color: rgb(232, 232, 232)")
        self.ui.process_pushButton3.setStyleSheet("background-color: rgb(232, 232, 232)")
        self.ui.lineEdit_BsKey.setFocus()
        config = configparser.ConfigParser()
        config.read("./config.ini",encoding="utf-8-sig")    ##utf-8-sig  & UTF-8
        self.Lotus_server = config['DEFAULT']['Lotus_server']
        self.outputfile = config['DEFAULT']['輸出路徑']
        fontcolor = config['DEFAULT']['忽略字休顏色']
        self.fontcolor = [int(i) for i in fontcolor.split(';')]
        # print("outputfile:", outputfile)
        # self.ui.comboBox1.setCurrentText(config['DEFAULT']['編號次序'])
        # self.ui.comboBox2.setCurrentText(config['DEFAULT']['PDF输出方向'])
        # self.ui.comboBox3.setCurrentText(config['DEFAULT']['出血位'])
        self.ui.lineEdit_r1c1.setText(config['DEFAULT']['橫排個數'])
        self.ui.lineEdit_r1c2.setText(config['DEFAULT']['橫刀位(寸)'])
        self.ui.lineEdit_r1c3.setText(config['DEFAULT']['左右出血位(寸)'])
        self.ui.lineEdit_r1c4.setText(config['DEFAULT']['上下出血位(寸)'])
        self.ui.lineEdit_r1c5.setText(config['DEFAULT']['右邊增減'])
        self.ui.lineEdit_r1c6.setText(config['DEFAULT']['下邊增減'])
        self.ui.lineEdit_r2c1.setText(config['DEFAULT']['豎排個數'])
        self.ui.lineEdit_r2c2.setText(config['DEFAULT']['豎刀位(寸)'])
        # self.ui.lineEdit_r2c3.setText(config['DEFAULT']['實際成品長'])
        # self.ui.lineEdit_r2c4.setText(config['DEFAULT']['實際成品寬'])
        self.ui.lineEdit_r3c3.setText(config['DEFAULT']['水平位置'])
        self.ui.lineEdit_r3c4.setText(config['DEFAULT']['垂直位置'])
        self.ui.lineEdit_r3c1.setText(config['DEFAULT']['一箱張數'])
        self.ui.lineEdit_r3c2.setText(config['DEFAULT']['一個PDF檔頁數'])
        # self.ui.frame_act_size.hide()
        # self.ui.frame_bleeding.hide()


    def revised_pdf_size(self):
        self.ui.frame_act_size.show()
        pdf_size_change=True

    def revised_bleeding(self):
        self.ui.frame_bleeding.show()


    def check_divmod(self):
        # if divmod(int(self.ui.lineEdit_r3c2.text()), int(self.ui.lineEdit_r3c1.text()))[1] != 0:
        #     self.ui.add_pushButton2.setStyleSheet("background-color: rgb(232, 232, 232)")
        #     self.ui.add_pushButton2.setEnabled(False)
        # else:
        #     self.ui.add_pushButton2.setStyleSheet("background-color: rgb(239, 239, 119)")
        #     self.ui.add_pushButton2.setEnabled(True)
        pass
    def check_bleeding(self):
        # if self.ui.comboBox3.currentIndex()>0:
        #     self.ui.lineEdit_r1c3.setEnabled(True)
        #     self.ui.lineEdit_r1c4.setEnabled(True)
        # else:
        #     self.ui.lineEdit_r1c3.setEnabled(False)
        #     self.ui.lineEdit_r1c4.setEnabled(False)
        pass

    def pdf_info(self):
        row = self.ui.open_tableView1.currentIndex().row()
        # colum = self.ui.open_tableView1.currentIndex().column()
        text1 = self.ui.open_tableView1.model().index(row, 1).data()
        text2 = self.ui.open_tableView1.model().index(row, 2).data()
        self.ui.lineEdit_r2c3.setText(text1)
        self.ui.lineEdit_r2c4.setText(text2)

        # print(text1,":",text2)



    def bs_PDF(self, PDFfilepath, start_seq,bskey,double):
        """
        處理鋪數PDF
        :return:
        """
        scanfold, select_filename = os.path.split(PDFfilepath)
        PDF_filename = os.path.splitext(select_filename)[0]
        scanfile = os.walk(scanfold)

        t1 = []
        for s in os.listdir(scanfold):
            if s.startswith(PDF_filename) and s == select_filename:
                t2 = (1, s)
                t1.append(t2)
            if s.startswith(PDF_filename) and s != select_filename:
                num1 = int(s[len(PDF_filename) + 4:len(s) - 4])
                t2 = (num1, s)
                t1.append(t2)

        t0 = sorted(t1, key=lambda x: x[0])
        # print(t0)
        doc0 = fitz.open(scanfold + "\\" + t0[0][1])
        width, height = doc0[0].mediabox_size
        # width, height = doc0[0].rect.width, doc0[0].rect.height
        r1 = fitz.Rect(0, 0, width, height)

        doc1 = fitz.open(scanfold + "\\" + t0[-1][1])
        std_qty = doc0.page_count
        end_qty = doc1.page_count
        if double:
            std_qty = int(std_qty / 2)
            end_qty = int(end_qty / 2)
        filelist = []
        for key, value in enumerate(t0):
            seq1 = start_seq + std_qty * key
            if key == len(t0) - 1:
                seq2 = start_seq + std_qty * key + (end_qty - 1)
            else:
                seq2 = start_seq + std_qty * key + (std_qty - 1)
            seq = (seq1, seq2)
            filelist.append(list(value + seq))

        print(filelist)
        seqlist = self.seqlist
        bs_list = []
        seqf = []
        seqpn = []
        for fl in filelist:
            # print(fl[2],fl[3])
            for seqnum in seqlist:
                if seqnum in range(fl[2], fl[3] + 1):
                    pagenum = divmod(seqnum, std_qty)[1]
                    if pagenum == 0:
                        pagenum = std_qty
                    seqf.append(seqnum)
                    if double:
                        pagenum = pagenum * 2 - 1
                        if pagenum == 0:
                            pagenum = std_qty * 2 - 1
                        seqpn.append(int(pagenum))
                        pagenum = pagenum + 1
                    seqpn.append(pagenum)
            fl.append(seqf)
            fl.append(seqpn)
            bs_list.append(fl)
            seqf = []
            seqpn = []

        print(bs_list)
        bs_list1 = [i for i in bs_list if i[4] != []]
        print(bs_list1)

        bs_doc = fitz.open()
        for bsfile in bs_list1:
            doc_seq = fitz.open(scanfold + "\\" + bsfile[1])
            bspn = bsfile[5]  # 補數PDF頁碼
            for pn in bspn:
                page = bs_doc.new_page(-1, width=width, height=height)
                page.show_pdf_page(r1, doc_seq, pn - 1, overlay=False)

        if bs_doc.page_count>0:
            bs_doc.save(self.outputfile+bskey+".pdf", garbage=0, deflate=True)  # 抽出了需要補數的單個PDF檔.
            return True
        else:
            QMessageBox.information(self, "錯誤提示!", "找不到補數PDF檔編號! 請檢查打開的PDF檔或者開始編號等問題")
            return False

    def bs_PDF_bak20230324(self, PDFfilepath, start_seq,bskey):
        """
        處理鋪數PDF
        :return:
        """
        scanfold, select_filename = os.path.split(PDFfilepath)
        PDF_filename = os.path.splitext(select_filename)[0]
        scanfile = os.walk(scanfold)

        t1 = []
        for s in os.listdir(scanfold):
            if s.startswith(PDF_filename) and s == select_filename:
                t2 = (1, s)
                t1.append(t2)
            if s.startswith(PDF_filename) and s != select_filename:
                num1 = int(s[len(PDF_filename) + 4:len(s) - 4])
                t2 = (num1, s)
                t1.append(t2)

        t0 = sorted(t1, key=lambda x: x[0])
        # print(t0)
        doc0 = fitz.open(scanfold + "\\" + t0[0][1])
        width, height = doc0[0].mediabox_size
        # width, height = doc0[0].rect.width, doc0[0].rect.height
        r1 = fitz.Rect(0, 0, width, height)

        doc1 = fitz.open(scanfold + "\\" + t0[-1][1])
        std_qty = doc0.page_count
        end_qty = doc1.page_count

        filelist = []
        for key, value in enumerate(t0):
            seq1 = start_seq + std_qty * key
            if key == len(t0) - 1:
                seq2 = start_seq + std_qty * key + (end_qty - 1)
            else:
                seq2 = start_seq + std_qty * key + (std_qty - 1)
            seq = (seq1, seq2)
            filelist.append(list(value + seq))

        print(filelist)
        seqlist = self.seqlist
        bs_list = []
        seqf = []
        seqpn = []
        for fl in filelist:
            # print(fl[2],fl[3])
            for seqnum in seqlist:
                if seqnum in range(fl[2], fl[3] + 1):
                    pagenum = divmod(seqnum, std_qty)[1]
                    if pagenum == 0:
                        pagenum = std_qty
                    seqf.append(seqnum)
                    seqpn.append(pagenum)
            fl.append(seqf)
            fl.append(seqpn)
            bs_list.append(fl)
            seqf = []
            seqpn = []

        print(bs_list)
        bs_list1 = [i for i in bs_list if i[4] != []]
        print(bs_list1)

        bs_doc = fitz.open()
        for bsfile in bs_list1:
            doc_seq = fitz.open(scanfold + "\\" + bsfile[1])
            bspn = bsfile[5]  # 補數PDF頁碼
            for pn in bspn:
                page = bs_doc.new_page(-1, width=width, height=height)
                page.show_pdf_page(r1, doc_seq, pn - 1, overlay=False)

        bs_doc.save(self.outputfile+bskey+".pdf", garbage=0, deflate=True)  # 抽出了需要補數的單個PDF檔.

    def bs_PDF_group(self, PDFfilepath, start_seq,bskey):
        """
        處理鋪數PDF
        :return:
        """
        scanfold, select_filename = os.path.split(PDFfilepath)
        PDF_filename = os.path.splitext(select_filename)[0]
        scanfile = os.walk(scanfold)

        t1 = []
        for s in os.listdir(scanfold):
            if s.startswith(PDF_filename) and s == select_filename:
                t2 = (1, s)
                t1.append(t2)
            if s.startswith(PDF_filename) and s != select_filename:
                num1 = int(s[len(PDF_filename) + 4:len(s) - 4])
                t2 = (num1, s)
                t1.append(t2)

        t0 = sorted(t1, key=lambda x: x[0])
        # print(t0)
        doc0 = fitz.open(scanfold + "\\" + t0[0][1])
        width, height = doc0[0].mediabox_size
        # width, height = doc0[0].rect.width, doc0[0].rect.height
        r1 = fitz.Rect(0, 0, width, height)

        doc1 = fitz.open(scanfold + "\\" + t0[-1][1])
        std_qty = doc0.page_count
        end_qty = doc1.page_count

        filelist = []
        for key, value in enumerate(t0):
            seq1 = start_seq + std_qty * key
            if key == len(t0) - 1:
                seq2 = start_seq + std_qty * key + (end_qty - 1)
            else:
                seq2 = start_seq + std_qty * key + (std_qty - 1)
            seq = (seq1, seq2)
            filelist.append(list(value + seq))

        print(filelist)
        seqlist = self.seqlist
        bs_list1 = []
        for seqnum in seqlist:
            for fl in filelist:
                if seqnum in range(fl[2], fl[3] + 1):
                    pagenum = divmod(seqnum, std_qty)[1]
                    if pagenum == 0:
                        pagenum = std_qty
                    bs_list1.append((fl[0],fl[1],fl[2],fl[3],seqnum,pagenum))
        print(bs_list1)

        bs_doc = fitz.open()
        for bsfile in bs_list1:
            doc_seq = fitz.open(scanfold + "\\" + bsfile[1])
            bspn = bsfile[5]  # 補數PDF頁碼

            page = bs_doc.new_page(-1, width=width, height=height)
            page.show_pdf_page(r1, doc_seq, bspn - 1, overlay=False)

        if bs_doc.page_count>0:
            bs_doc.save(self.outputfile+bskey+".pdf", garbage=0, deflate=True)  # 抽出了需要補數的單個PDF檔.
            return True
        else:
            QMessageBox.information(self, "錯誤提示!", "找不到補數PDF檔編號! 請檢查打開的PDF檔或者開始編號等問題")
            return False

    def add_File(self):
        print("isChecked:",self.ui.checkBox_blank.isChecked())
        self.ui.add_pushButton2.setStyleSheet("background-color: rgb(232, 232, 232)")
        if self.ui.checkBox_blank.isChecked() == True and self.ui.label_art.text() == 'no_art':
            QMessageBox.information(self, "錯誤提示!", "加隔紙需要有稿件!")
            return
        self.ui.add_pushButton2.setEnabled(False)
        # self.ui.frame_bleeding.hide()
        # self.ui.frame_act_size.hide()
        tableView1_model=self.ui.open_tableView1.model()
        filepath = tableView1_model.index(0, 5).data()
        start_seq = int(self.ui.lineEdit_StartSeq.text())



        if len(self.BsKey)!=0:
            bskey=self.BsKey
        else:
            bskey=self.ui.lineEdit_BsKey.text()
        double=self.ui.checkBox1.isChecked()
        if self.group:
            checkPDF1=self.bs_PDF_group(filepath, start_seq, bskey)
        else:
            checkPDF1=self.bs_PDF(filepath,start_seq,bskey,double)
        # print("check",self.ui.process_tableView2.model())
        # print(mode1.rowCount())
        if self.ui.process_tableView2.model() is None:
            view2_rowCount=0

        else:
            view2_rowCount=self.ui.process_tableView2.model().rowCount()

        if checkPDF1:   #如果有PDF檔,
            # print("view2_rowCount=",view2_rowCount, " range:",view2_rowCount+tableView1_model.rowCount())
            for i in range(view2_rowCount,view2_rowCount+tableView1_model.rowCount()):
                filename=bskey+".pdf"

                page_count = self.seqlist.__len__()
                filepath = self.outputfile+bskey+".pdf"
                print("i=",i,"filename:",filename)
                model_view2.setItem(i, 0, QStandardItem(filename))
                model_view2.setItem(i, 1, QStandardItem(self.ui.comboBox1.currentText()))
                model_view2.setItem(i, 2, QStandardItem(self.ui.comboBox2.currentText()))
                model_view2.setItem(i, 3, QStandardItem(self.ui.comboBox3.currentText()))
                model_view2.setItem(i, 4, QStandardItem(self.ui.lineEdit_r1c1.text()))
                model_view2.setItem(i, 5, QStandardItem(self.ui.lineEdit_r1c2.text()))
                model_view2.setItem(i, 6, QStandardItem(self.ui.lineEdit_r1c3.text()))
                model_view2.setItem(i, 7, QStandardItem(self.ui.lineEdit_r1c4.text()))
                model_view2.setItem(i, 8, QStandardItem(self.ui.lineEdit_r2c1.text()))
                model_view2.setItem(i, 9, QStandardItem(self.ui.lineEdit_r2c2.text()))
                model_view2.setItem(i, 10, QStandardItem(self.ui.lineEdit_r2c3.text()))
                model_view2.setItem(i, 11, QStandardItem(self.ui.lineEdit_r2c4.text()))
                model_view2.setItem(i, 12, QStandardItem(self.ui.lineEdit_r3c1.text()))
                model_view2.setItem(i, 13, QStandardItem(self.ui.lineEdit_r3c2.text()))
                model_view2.setItem(i, 14, QStandardItem(page_count))   #总页数
                model_view2.setItem(i, 15, QStandardItem(filepath)) #檔案路徑
                model_view2.setItem(i, 16, QStandardItem(self.ui.lineEdit_r3c3.text()))     #水平位置
                model_view2.setItem(i, 17, QStandardItem(self.ui.lineEdit_r3c4.text()))     #垂直位置
                model_view2.setItem(i, 18, QStandardItem(self.ui.checkBox1.isChecked().__str__()))       #是否雙面
                model_view2.setItem(i, 19, QStandardItem(self.ui.lineEdit_r1c5.text()))     #右邊增減
                model_view2.setItem(i, 20, QStandardItem(self.ui.lineEdit_r1c6.text()))     #下邊增減
                model_view2.setItem(i, 21, QStandardItem(self.ui.label_art.text()))         #排版PDF稿件
                model_view2.setItem(i, 22, QStandardItem(self.ui.checkBox_blank.isChecked().__str__()))     #加隔紙
                model_view2.setItem(i, 23, QStandardItem(self.ui.checkBox_blank2.isChecked().__str__()))    #單面加稿件底頁
                model_view2.setItem(i, 24, QStandardItem(self.ui.comboBox_rotate.currentText()))        #PDF角度
                model_view2.setItem(i, 25, QStandardItem(self.ui.checkBox_addpage.isChecked().__str__()))       #加頁碼
                model_view2.setItem(i, 26, QStandardItem(self.ui.lineEdit_Repeat.text()))       #重復數量

            self.ui.process_tableView2.setModel(model_view2)
            # self.ui.process_tableView2.resizeColumnToContents(0)
            self.ui.process_tableView2.resizeColumnsToContents()
            self.ui.process_tableView2.resizeRowsToContents()
            if model_view2.rowCount()>0:
                self.ui.process_pushButton3.setStyleSheet("background-color: rgb(239, 239, 119)")
                self.ui.process_pushButton3.setEnabled(True)

            self.ui.open_tableView1.model().removeRows(0,self.ui.open_tableView1.model().rowCount())
            # self.ui.label_art.setText("no_art")

    def batch_combine(self):
        self.ui.process_pushButton3.setStyleSheet("background-color: rgb(232, 232, 232)")
        self.ui.process_pushButton3.setEnabled(False)
        self.ui.open_pushButton1.setStyleSheet("background-color: rgb(232, 232, 232)")
        self.ui.open_pushButton1.setEnabled(False)
        self.ui.open_pushButton1.hide()
        print(self.ui.open_pushButton1.isEnabled())
        tableView2_model=self.ui.process_tableView2.model()
        self.ui.progressBar.setRange(0, tableView2_model.rowCount())    #設置進度條
        tableView2_list=[]
        for i in range(tableView2_model.rowCount()):
            if not os.path.isdir(outputfile+str(tableView2_model.index(i, 0).data()).split('.')[0]):
                pass
                # os.mkdir(outputfile+str(tableView2_model.index(i, 0).data()).split('.')[0])
            print("folder:", outputfile+str(tableView2_model.index(i, 0).data()).split('.')[0])
            for j in range(tableView2_model.columnCount()):
                tableView2_list.append(tableView2_model.index(i, j).data())

            combine(tableView2_list)
            self.ui.progressBar.setValue(i+1)
            tableView2_list.clear()
        self.ui.process_tableView2.model().removeRows(0,self.ui.process_tableView2.model().rowCount())

        self.ui.open_pushButton1.setStyleSheet("background-color: rgb(239, 239, 119)")
        self.ui.open_pushButton1.setEnabled(True)
        self.ui.open_pushButton1.show()

    def get_text_color(self,page,num):
        color_list=[]
        find_seq = []
        blocks = page.get_text("dict", flags=11)["blocks"]
        for b in blocks:  # iterate through the text blocks
            for l in b["lines"]:  # iterate through the text lines

                # 查找编号
                keyword = self.seqcode + str(self.start_seq + num).strip()
                # print("keyword", keyword, 'text', s['text'].upper())
                data = l['spans']
                for lines in data:
                    # print("lines",lines)
                    if keyword == lines['text'].upper():  # only store font information of a specific keyword
                        find_seq.append((lines['text'], lines['size'], lines['font']))
                        print(lines['text'], lines['size'], lines['font'])
                # 查找编号

                for s in l["spans"]:  # iterate through the text spans
                    checkjob = re.findall(r'\d{7}-\d{1}.*\.(pdf|indd|eps)   \d{1,7}', s["text"], re.I)
                    checkdatetime = re.findall(r'\d{1,4}/\d{1,2}/\d{1,4}   \d{1,2}:\d{1,2}', s["text"], re.I)


                    # print(s["color"],type(s["color"]))
                    #2236191=CMYK 黑色(或直接插入PDF稿件的文字), 0=RGB 黑色; 16777215=白色
                    if (s["color"] not in self.fontcolor) and not (checkjob or checkdatetime):
                    # if (s["color"] != 0 and s["color"] != 2236191) and not (checkjob or checkdatetime):
                        if s["color"] == 16777215:
                            color = '反白文字('+str(s["color"])+'):'
                        else:
                            color = '顏色文字('+str(s["color"])+'):'
                        color_list.append(color + s["text"])
        print('find_seq:',find_seq)
        return color_list,find_seq

    def get_pdf_info(self, filelist):
        model = QStandardItemModel()
        model.setHorizontalHeaderLabels(["工單版本","文件名", "尺寸-寬", "尺寸-高","PDF頁數", "路徑"])
        for row_num,fl in enumerate(filelist):
            # print("check _job",fl+str(re.search('_job', fl) is not None))
            if re.search('_job', fl) is not None:
                QMessageBox.information(ui_mainwindow,"錯誤提示!","不能選擇含有'_job'的文件"+chr(13)+fl)
            else:
                src = fitz.open(fl)
                page_count_last = src.page_count
                width, height = src[0].mediabox_size

                doc0=src[0]
                # color_list01 = []
                color_list01,find_seq0=self.get_text_color(doc0,0)
                # width, height = src[0].rect.width, src[0].rect.height
                #字體Embedded提示
                fontslist = src[0].get_fonts()
                if page_count_last > 1:
                    doc1 = src[1]
                    color_list1,find_seq1 = self.get_text_color(doc1,1)


                    for cl1 in color_list1:
                        color_list01.append(cl1)
                    fontslist2 = src[1].get_fonts()
                    for emb2 in fontslist2:
                        fontslist.append(emb2)
                for Emb in fontslist:
                    if Emb[1] == 'n/a':
                        QMessageBox.information(ui_mainwindow, "字體嵌入提示!", fl+chr(10)+Emb[3] + ": 不能嵌入字體")

                if color_list01:
                    show_text=''
                    for cl in sorted(color_list01,reverse=False):
                        show_text=show_text+cl+chr(10)
                    with open('PDF_font_color.txt', 'w',encoding="UTF-8") as f:
                        f.write(show_text)
                    QMessageBox.information(ui_mainwindow, "打印有顏色的字體補數提示!", show_text)

                #字體Embedded提示
                width_inch = str(width/one_inch)
                height_inch = str(height / one_inch)
                filename=str(fl).split("/")[-1]
                if filename[0]!='1':
                    fx = str(filename).split(" ")
                    Hjob = "1" + fx[0].replace("-", "")[2:8]
                    Hlot = fx[1][1:]
                    Hjoblot = Hjob+' L'+Hlot
                else:
                    Hjoblot=filename[0:7]
                model.setItem(row_num, 0, QStandardItem(Hjoblot))
                model.setItem(row_num, 1, QStandardItem(str(fl).split("/")[-1]))
                model.setItem(row_num, 2, QStandardItem(width_inch))
                model.setItem(row_num, 3, QStandardItem(height_inch))
                model.setItem(row_num, 4, QStandardItem(str(page_count_last)))
                model.setItem(row_num, 5, QStandardItem(fl))


        return model

    def Open_File(self):
        global dirpath
        self.ui.progressBar.setValue(0)
        filelist, _ = QFileDialog.getOpenFileNames(
            self,  # 父窗口对象
            "选择要處理的PDF檔",  # 标题
            # r"./",  # 起始目录
            dirpath,
            "數據類型 (*.PDF)"  # 选择类型过滤项，过滤内容在括号中
        )

        # table_name = self.ui.listWidget_data_list.currentItem().whatsThis()

        if filelist!=[]:
            print("filelist:",filelist)
            print("path:",os.path.split(filelist[0])[0]+"/")
            dirpath=os.path.split(filelist[0])[0]+"/"
            model = self.get_pdf_info(filelist)
            self.ui.open_tableView1.setModel(model)
            self.ui.open_tableView1.resizeColumnToContents(0)
            self.ui.open_tableView1.resizeColumnToContents(1)
            self.ui.open_tableView1.resizeRowsToContents()
            self.ui.lineEdit_r2c3.setText(self.ui.open_tableView1.model().index(0, 2).data())
            self.ui.lineEdit_r2c4.setText(self.ui.open_tableView1.model().index(0, 3).data())
            self.ui.lineEdit_r3c3.setText("0")
            self.ui.lineEdit_r3c4.setText("0")
            self.ui.lineEdit_r1c5.setText("0")
            self.ui.lineEdit_r1c6.setText("0")

            self.ui.lineEdit_r1c1.setText("1")
            self.ui.lineEdit_r1c2.setText("0")
            self.ui.lineEdit_r1c3.setText("0")
            self.ui.lineEdit_r1c4.setText("0")
            self.ui.lineEdit_r2c1.setText("1")
            self.ui.lineEdit_r2c2.setText("0")
            self.ui.lineEdit_r3c1.setText("1000")
            self.ui.lineEdit_r3c2.setText("2000")
            self.ui.comboBox1.setCurrentText('3:兜圈順序打印')
            self.ui.checkBox1.setChecked(False)
            self.ui.checkBox_blank.setChecked(False)
            self.ui.checkBox_blank2.setChecked(False)
            self.ui.checkBox_addpage.setChecked(False)


            self.ui.add_pushButton2.setStyleSheet("background-color: rgb(239, 239, 119)")
            self.ui.add_pushButton2.setEnabled(True)
        else:
            print("文件名空")

    def Open_File2(self):
        global dirpath
        self.ui.progressBar.setValue(0)
        filelist, _ = QFileDialog.getOpenFileNames(
            self,  # 父窗口对象
            "选择要處理的PDF檔",  # 标题
            # r"./",  # 起始目录
            dirpath,
            "數據類型 (*.PDF)"  # 选择类型过滤项，过滤内容在括号中
        )

        # table_name = self.ui.listWidget_data_list.currentItem().whatsThis()

        if filelist!=[]:
            self.start_seq = int(self.ui.lineEdit_StartSeq.text())

            print("filelist:",filelist)
            print("path:",os.path.split(filelist[0])[0]+"/")
            dirpath=os.path.split(filelist[0])[0]+"/"
            model = self.get_pdf_info(filelist)
            self.ui.open_tableView1.setModel(model)
            self.ui.open_tableView1.resizeColumnToContents(0)
            self.ui.open_tableView1.resizeColumnToContents(1)
            self.ui.open_tableView1.resizeRowsToContents()
            self.ui.lineEdit_r2c3.setText(self.ui.open_tableView1.model().index(0, 2).data())
            self.ui.lineEdit_r2c4.setText(self.ui.open_tableView1.model().index(0, 3).data())
            # self.ui.lineEdit_r3c3.setText("0")
            # self.ui.lineEdit_r3c4.setText("0")
            # self.ui.lineEdit_r1c5.setText("0")
            # self.ui.lineEdit_r1c6.setText("0")
            #
            # self.ui.lineEdit_r1c1.setText("1")
            # self.ui.lineEdit_r1c2.setText("0")
            # self.ui.lineEdit_r1c3.setText("0")
            # self.ui.lineEdit_r1c4.setText("0")
            # self.ui.lineEdit_r2c1.setText("1")
            # self.ui.lineEdit_r2c2.setText("0")
            # self.ui.lineEdit_r3c1.setText("1000")
            # self.ui.lineEdit_r3c2.setText("2000")
            self.ui.comboBox1.setCurrentText('1:從小到大')
            self.ui.checkBox1.setChecked(False)
            self.ui.checkBox_blank.setChecked(False)
            self.ui.checkBox_blank2.setChecked(False)
            self.ui.checkBox_addpage.setChecked(False)


            self.ui.add_pushButton2.setStyleSheet("background-color: rgb(239, 239, 119)")
            self.ui.add_pushButton2.setEnabled(True)
        else:
            print("文件名空")

    def Open_artFile(self):
        self.ui.progressBar.setValue(0)

        artfile, _ = QFileDialog.getOpenFileNames(
            self,  # 父窗口对象
            "选择稿件PDF檔",  # 标题
            dirpath,  # 起始目录
            "數據類型 (*.PDF)"  # 选择类型过滤项，过滤内容在括号中
        )
        if artfile != []:
            self.ui.label_art.setText(artfile[0])
        else:
            print("稿件沒有選擇!")
        # table_name = self.ui.listWidget_data_list.currentItem().whatsThis()

def sort_data1(std_count, end_count, Layout, cnt_page, std_pdf_file_count2):
    """
    順序一棟落Data 處理.
    :param std_count:一個PDF檔記錄數量
    :param end_count:尾數箱剩余數量.
    :param Layout: 排版個數
    :param cnt_page:一箱張數.
    :param std_pdf_file_count2:最后一個PDF檔里有多少個是標準張數(一箱張數).
    :return: list1, list2_start, list2_end, cnt_page_end
    list1: (seq, cnt_number, layout_seq, pdf_page, layout_number):  正常一個PDF檔的列表
        seq: 編號(順序號)
        cnt_number: 箱號
        layout_seq:排版順序號
        pdf_page: PDF 頁碼	**有用到L[3]
        layout_number: 排版號 **有用到L[4]

    list2_start : 最后一個檔案的標準箱列表.
    list2_end: 最后一個檔案的尾箱列表.
    cnt_page_end: 尾箱頁數
    """
    # 順序一棟落Data 處理.
    # 參數
    # data_count = 12125  # data 數量
    # width_layout = 3  # 寬排版個數
    # height_layout = 2  # 高排版個數
    # cnt_page = 100  # 一箱張數
    # std_pdf_file_count2 尾箱整箱有多少箱
    # layout = width_layout * height_layout  # 排版數
    cnt_qty = cnt_page * Layout  # 一箱數量
    end_cnt_page = std_pdf_file_count2 * cnt_qty

    list1 = []  # 一棟落順序
    for seq in range(1, std_count + 1):  # seq  編號
        cnt_number = math.ceil(seq / cnt_qty)  # cnt_number 箱號
        layout_seq = math.ceil(seq / cnt_page)  # layout_seq 排版順序號
        pdf_page = (divmod(seq - 1, cnt_page)[1] + 1) + cnt_number * cnt_page - cnt_page  # pdf_page   PDF 頁碼
        layout_number = divmod(layout_seq - 1, Layout)[1] + 1  # layout_number排版號,  1,2,3,1,2,3,1,2,3
        t1 = (seq, cnt_number, layout_seq, pdf_page, layout_number)
        list1.append(t1)

    # list2_start = list1[0:end_cnt_page]
    list2_start = list1[0:end_cnt_page]
    list2_end = []

    cnt_page_end = math.ceil(end_count / Layout)
    cnt_qty_end = cnt_page_end * Layout
    for seq in range(1, end_count + 1):  # seq  編號
        cnt_number_end = math.ceil(seq / cnt_qty_end)  # cnt_number 箱號
        layout_seq_end = math.ceil(seq / cnt_page_end)  # layout_seq 排版順序號
        pdf_page_end = (divmod(seq - 1, cnt_page_end)[1] + 1)  # pdf_page   PDF 頁碼
        layout_number = divmod(layout_seq_end - 1, Layout)[1] + 1  # layout_number排版號,  1,2,3,1,2,3,1,2,3
        # t2=(std_count+seq, cnt_number_end, layout_seq_end, std_page+pdf_page_end, layout_number)
        t2 = (seq, cnt_number_end, layout_seq_end, pdf_page_end, layout_number)
        list2_end.append(t2)

    return list1, list2_start, list2_end, cnt_page_end

def sort_data1_double(std_pdf_file_count1, end_count, width_layout, height_layout, cnt_page, std_pdf_file_count2,pdf_file_page_qty):
    cntqty=int(pdf_file_page_qty/cnt_page)  #一個檔案有多少箱
    Layout = width_layout * height_layout
    cnt_qty = cnt_page * Layout  # 一箱數量
    end_cnt_page = std_pdf_file_count2 * cnt_qty  # 尾數PDF檔整箱數量
    list1 = []  # 一棟落順序
    for cnt in range(cntqty):
        for h in range(1, height_layout + 1):
            for w in range(1, width_layout + 1):
                for pdf_page in range(1, cnt_page + 1):
                    m2 = divmod(pdf_page - 1, 2)[1] + 1  # 2是雙面里的底面
                    m3 = divmod(w - 1, width_layout)[1] + 1  # 判斷橫排是否在第一個或最后一個
                    p1 = (h - 1) * width_layout + w  # 第一頁的標準layout
                    if m2 == 2 and m3 == 1:
                        layout_number = p1 + width_layout - 1
                    elif m2 == 2 and m3 == width_layout:
                        layout_number = p1 - width_layout + 1
                    else:
                        layout_number = p1
                    page_order=pdf_page + (cnt_page * cnt)  #順序
                    t1 = (0, 0, 0, page_order, layout_number)
                    list1.append(t1)
                    # print("ta:", t1)
                    # print("頁碼a:", pdf_page + (cnt_page * cnt), "lay:", layout_number, "橫排:", w, "豎排:", h, "p1:", p1,
                    #       "m2:", m2, "m3:", m3, "b  - a * (height_layout): ", w - h * (height_layout))
    list2_start = list1[0:end_cnt_page]
    list2_end = []
    cnt_page_end = math.ceil(end_count / Layout)

    while divmod(cnt_page_end,Layout)[1]>0:    #如果尾數箱不能整除Layout, 就需要加至能整除.
        cnt_page_end=cnt_page_end+1
    # for cnt in range(std_pdf_file_count2):
    for cnt in range(1):
        for h in range(1, height_layout + 1):
            for w in range(1, width_layout + 1):
                for pdf_page in range(1, cnt_page_end + 1):
                    m2 = divmod(pdf_page - 1, 2)[1] + 1  # 2是雙面里的底面
                    m3 = divmod(w - 1, width_layout)[1] + 1  # 判斷橫排是否在第一個或最后一個
                    p1 = (h - 1) * width_layout + w  # 第一頁的標準layout
                    if m2 == 2 and m3 == 1:
                        layout_number = p1 + width_layout - 1

                    elif m2 == 2 and m3 == width_layout:
                        layout_number = p1 - width_layout + 1

                        # print("test-layout_number:",p1,width_layout,layout_number)

                    else:
                        layout_number = p1
                    page_order = pdf_page + (cnt_page_end * cnt)    #順序
                    # page_order = pdf_page
                    t1 = (0, 0, 0, page_order, layout_number)
                    # print("page_order:",page_order)
                    # print("tb:", t1)
                    list2_end.append(t1)
                    # print("頁碼b:",pdf_page+(cnt_page*cnt), "lay:",layout_number, "橫排:",w, "豎排:",h,"p1:",p1,"m2:",m2, "m3:",m3,"b  - a * (height_layout): ",w  - h * (height_layout))
    return list1, list2_start, list2_end, cnt_page_end

def sort_data1_fx1(width_layout, height_layout, end_count, cnt_page, pdf_file_page_qty,std_pdf_file_count2):
    """
    分切成品 (單個) - 順序 處理
    :return:
    """
    # end_count = 59  #尾箱數量
    # width_layout = 2  # 寬排版個數
    # height_layout = 3  # 高排版個數
    # cnt_page = 100  # 一箱張數
    # pdf_file_page_qty=200
    # std_pdf_file_count2=1     #尾箱整箱有多少箱

    cnts=int(pdf_file_page_qty / cnt_page)  #一個PDF檔有多少箱

    layout = width_layout * height_layout  # 排版數
    cnt_qty = cnt_page * layout  # 一箱數量
    end_cnt_page = std_pdf_file_count2 * cnt_qty

    # 標準箱數據
    list1 = []  # 一棟落順序
    pageseq = 0
    pagenum = 0
    add1 = 0
    for cnt in range(1, cnts + 1):
        layout_number = 0
        pagenum = 0
        add1 = 0
        for w_lay in range(1, width_layout + 1):
            pagenum = 0
            add1 =(w_lay-1)*height_layout
            for page in range(1, cnt_page + 1):
                pagenum = page + (cnt - 1) * cnt_page
                for h_lay in range(1, height_layout+1):
                    layout_number=divmod(h_lay - 1, height_layout)[1] + 1 + add1
                    t1 = (0, 0, 0, pagenum, layout_number)
                    list1.append(t1)
                    # print(w_lay,pagenum,layout_number)


    #尾箱整箱數據
    list2_start = list1[0:end_cnt_page]

    #尾箱數據
    cnt = 1  # 尾箱只有一箱
    list2_end = []
    cnt_page_end = math.ceil(end_count / layout)  # 尾箱有多少張
    while divmod(cnt_page_end,layout)[1]>0:    #如果尾數箱不能整除Layout, 就需要加至能整除.
        cnt_page_end=cnt_page_end+1

    for w_lay in range(1, width_layout + 1):
        pagenum = 0
        add1 = (w_lay - 1) * height_layout
        for page in range(1, cnt_page_end + 1):
            pagenum = page + (cnt - 1) * cnt_page_end
            for h_lay in range(1, height_layout + 1):
                layout_number = divmod(h_lay - 1, height_layout)[1] + 1 + add1
                print(w_lay, pagenum, layout_number)
                t1 = (0, 0, 0, pagenum, layout_number)
                list2_end.append(t1)

    return list1, list2_start, list2_end, cnt_page_end

def sort_data1_fx2(width_layout, height_layout, end_count, cnt_page, pdf_file_page_qty,std_pdf_file_count2):
    """
    分切成品 (單個) - 倒序 處理
    :return:
    """
    # end_count = 59  #尾箱數量
    # width_layout = 2  # 寬排版個數
    # height_layout = 3  # 高排版個數
    # cnt_page = 100  # 一箱張數
    # pdf_file_page_qty=200
    # std_pdf_file_count2=1     #尾箱整箱有多少箱

    cnts=int(pdf_file_page_qty / cnt_page)  #一個PDF檔有多少箱

    layout = width_layout * height_layout  # 排版數
    cnt_qty = cnt_page * layout  # 一箱數量
    end_cnt_page = std_pdf_file_count2 * cnt_qty

    # 標準箱數據
    list1 = []  # 一棟落順序
    pageseq = 0
    pagenum = 0
    add1 = 0
    for cnt in range(1, cnts + 1):
        layout_number = 0
        pagenum = 0
        add1 = 0
        for w_lay in range(1, width_layout + 1):
            pagenum = 0
            add1 =(w_lay-1)*height_layout
            for page in range(cnt_page, 0 , -1):
                pagenum = page + (cnt - 1) * cnt_page
                for h_lay in range(height_layout, 0, -1):
                    layout_number=divmod(h_lay - 1, height_layout)[1] + 1 + add1
                    t1 = (0, 0, 0, pagenum, layout_number)
                    list1.append(t1)
                    print(w_lay,pagenum,layout_number)


    #尾箱整箱數據
    list2_start = list1[0:end_cnt_page]

    #尾箱數據
    cnt=1   #尾箱只有一箱
    list2_end = []
    cnt_page_end = math.ceil(end_count / layout)  # 尾箱有多少張
    for w_lay in range(1, width_layout + 1):
        pagenum = 0
        add1 = (w_lay - 1) * height_layout
        for page in range(cnt_page_end, 0 , -1):
            pagenum = page + (cnt - 1) * cnt_page_end
            for h_lay in range(height_layout, 0, -1):
                layout_number = divmod(h_lay - 1, height_layout)[1] + 1 + add1
                # print(w_lay, pagenum, layout_number)
                t1 = (0, 0, 0, pagenum, layout_number)
                list2_end.append(t1)

    return list1, list2_start, list2_end, cnt_page_end

def sort_data2(std_count, end_count, Layout, cnt_page, std_pdf_file_count2):
    # 倒序一棟落Data 處理.
    # 參數
    # data_count = 12125  # data 數量
    # width_layout = 3  # 寬排版個數
    # height_layout = 2  # 高排版個數
    # cnt_page = 100  # 一箱張數
    # std_pdf_file_count2 尾箱整箱有多少箱
    # layout = width_layout * height_layout  # 排版數
    cnt_qty = cnt_page * Layout  # 一箱數量
    # end_cnt_page = std_pdf_file_count2 * cnt_page
    end_cnt_page = std_pdf_file_count2 * cnt_qty
    list1 = []  # 一棟落順序
    for seq in range(1, std_count + 1):  # seq  編號
        cnt_number = math.ceil(seq / cnt_qty)  # cnt_number 箱號
        layout_seq = math.ceil(seq / cnt_page)  # layout_seq 排版順序號
        # pdf_page=(divmod(seq - 1, cnt_page)[1] + 1)+ cnt_number*cnt_page-cnt_page  #pdf_page   PDF 頁碼
        pdf_page = cnt_number * cnt_page - ((seq - 1) % cnt_page + 1) + 1  # pdf_page   PDF 頁碼
        layout_number = divmod(layout_seq - 1, Layout)[1] + 1  # layout_number排版號,  1,2,3,1,2,3,1,2,3
        t1 = (seq, cnt_number, layout_seq, pdf_page, layout_number)
        list1.append(t1)

    list2_start = list1[0:end_cnt_page]
    list2_end = []
    cnt_page_end = math.ceil(end_count / Layout)
    cnt_qty_end = cnt_page_end * Layout
    for seq in range(1, end_count + 1):  # seq  編號
        cnt_number_end = math.ceil(seq / cnt_qty_end)  # cnt_number 箱號
        layout_seq_end = math.ceil(seq / cnt_page_end)  # layout_seq 排版順序號
        # pdf_page_end=(divmod(seq - 1, cnt_page_end)[1] + 1)  #pdf_page   PDF 頁碼
        pdf_page_end = cnt_number_end * cnt_page_end - ((seq - 1) % cnt_page_end + 1) + 1  # pdf_page   PDF 頁碼
        layout_number = divmod(layout_seq_end - 1, Layout)[1] + 1  # layout_number排版號,  1,2,3,1,2,3,1,2,3
        # t2=(std_count+seq, cnt_number_end, layout_seq_end, std_page+pdf_page_end, layout_number)
        t2 = (seq, cnt_number_end, layout_seq_end, pdf_page_end, layout_number)
        list2_end.append(t2)

    return list1, list2_start, list2_end, cnt_page_end

def sort_data2_double(std_pdf_file_count1, end_count, width_layout, height_layout, cnt_page, std_pdf_file_count2,pdf_file_page_qty):
    cntqty = int(pdf_file_page_qty / cnt_page)  #一個檔案有多少箱
    Layout = width_layout * height_layout
    cnt_qty = cnt_page * Layout  # 一箱數量
    end_cnt_page = std_pdf_file_count2 * cnt_qty  # 尾數PDF檔整箱數量
    list1 = []  # 一棟落順序
    # for cnt in range(std_pdf_file_count1):
    for cnt in range(cntqty):
        for h in range(1, height_layout + 1):
            for w in range(1, width_layout + 1):
                for pdf_page in range(1, cnt_page + 1):
                    m2 = divmod(pdf_page - 1, 2)[1] + 1  # 2是雙面里的底面
                    m3 = divmod(w - 1, width_layout)[1] + 1  # 判斷橫排是否在第一個或最后一個
                    p1 = (h - 1) * width_layout + w  # 第一頁的標準layout
                    if m2 == 2 and m3 == 1:
                        layout_number = p1 + width_layout - 1
                    elif m2 == 2 and m3 == width_layout:
                        layout_number = p1 - width_layout + 1
                    else:
                        layout_number = p1
                    page_order = cnt_page * (cnt + 1) - pdf_page + 1    #倒序
                    if m2 == 1:
                        page_order2 = page_order - 1
                    else:
                        page_order2 = page_order + 1

                    t1 = (0, 0, 0, page_order2, layout_number)
                    list1.append(t1)
                    # print("頁碼:", pdf_page + (cnt_page * cnt), "lay:", layout_number, "橫排:", w, "豎排:", h, "p1:", p1,
                    #       "m2:", m2, "m3:", m3, "b  - a * (height_layout): ", w - h * (height_layout))
    list2_start = list1[0:end_cnt_page]
    list2_end = []
    cnt_page_end = math.ceil(end_count / Layout)

    while divmod(cnt_page_end,Layout)[1]>0:    #如果尾數箱不能整除Layout, 就需要加至能整除.
        cnt_page_end=cnt_page_end+1

    # for cnt in range(std_pdf_file_count2):
    for cnt in range(1):
        for h in range(1, height_layout + 1):
            for w in range(1, width_layout + 1):
                for pdf_page in range(1, cnt_page_end + 1):
                    m2 = divmod(pdf_page - 1, 2)[1] + 1  # 2是雙面里的底面
                    m3 = divmod(w - 1, width_layout)[1] + 1  # 判斷橫排是否在第一個或最后一個
                    p1 = (h - 1) * width_layout + w  # 第一頁的標準layout
                    if m2 == 2 and m3 == 1:
                        layout_number = p1 + width_layout - 1
                    elif m2 == 2 and m3 == width_layout:
                        layout_number = p1 - width_layout + 1
                    else:
                        layout_number = p1
                    page_order = cnt_page_end * (cnt + 1) - pdf_page + 1    #倒序
                    if m2 == 1:
                        page_order2 = page_order - 1
                    else:
                        page_order2 = page_order + 1

                    t1 = (0, 0, 0, page_order2, layout_number)
                    list2_end.append(t1)
                    # print("頁碼:",pdf_page+(cnt_page*cnt), "lay:",layout_number, "橫排:",w, "豎排:",h,"p1:",p1,"m2:",m2, "m3:",m3,"b  - a * (height_layout): ",w  - h * (height_layout))
    return list1, list2_start, list2_end, cnt_page_end

def sort_data3(std_count, end_count, Layout, cnt_page, std_pdf_file_count2):
    # 順序兜圈Data 處理.
    # 參數
    # data_count = 12125  # data 數量
    # width_layout = 3  # 寬排版個數
    # height_layout = 2  # 高排版個數
    # cnt_page = 100  # 一箱張數
    # std_pdf_file_count2 尾箱整箱有多少箱
    # layout = width_layout * height_layout  # 排版數
    cnt_qty = cnt_page * Layout  # 一箱數量
    end_cnt_page = std_pdf_file_count2 * cnt_qty

    list1 = []  # 一棟落順序
    for seq in range(1, std_count + 1):  # seq  編號
        cnt_number = math.ceil(seq / cnt_qty)  # cnt_number 箱號
        layout_seq = math.ceil(seq / cnt_page)  # layout_seq 排版順序號
        pdf_page = math.ceil(seq/Layout)  # pdf_page   PDF 頁碼
        layout_number = divmod(seq - 1, Layout)[1] + 1  # layout_number排版號,  1,2,3,1,2,3,1,2,3
        t1 = (seq, cnt_number, layout_seq, pdf_page, layout_number)
        list1.append(t1)

    list2_start = list1[0:end_cnt_page]
    list2_end = []
    cnt_page_end = math.ceil(end_count / Layout)
    cnt_qty_end = cnt_page_end * Layout
    for seq in range(1, end_count + 1):  # seq  編號
        cnt_number_end = math.ceil(seq / cnt_qty_end)  # cnt_number 箱號
        layout_seq_end = math.ceil(seq / cnt_page_end)  # layout_seq 排版順序號
        pdf_page_end = math.ceil(seq/Layout)  # pdf_page   PDF 頁碼
        layout_number = divmod(seq - 1, Layout)[1] + 1  # layout_number排版號,  1,2,3,1,2,3,1,2,3
        t2 = (seq, cnt_number_end, layout_seq_end, pdf_page_end, layout_number)
        list2_end.append(t2)

    return list1, list2_start, list2_end, cnt_page_end

def combine(list1):
    print("combine")
    def doutimes(doc_process):
        # 重復頁
        # times = 3  # 重復次數
        # one_inch = 72  # 1寸=72像素, 常量
        # width_layout = 1  # 橫排個數
        # height_layout = 2  # 豎排個數
        # bleeding_w = 0.25  # 左右出血位(寸)
        # bleeding_H = 0.25  # 上下出血位(寸)
        # interval_w = 0  # 橫刀位(寸)
        # interval_h = 0  # 豎刀位(寸)
        # bleeding_type = "2:左上角水平位置"  # 2:左上角垂直位置
        # doc = fitz.open("./testadd3.pdf")
        # docpage = fitz.open("./testadd3.pdf")
        doc = doc_process
        rect_height = 7  # 方框高度
        page_count = doc.page_count
        width, height = doc[0].mediabox_size
        # width, height = doc[0].rect.width, doc[0].rect.height
        if bleeding_type[0] == "1":
            layout = width_layout
            bleeding = bleeding_w
            Length = width
            rotate = 180
        else:
            layout = height_layout
            bleeding = bleeding_H
            Length = height
            rotate = 270

        space_Interval = int(((Length - (bleeding * 2) / layout) - 10 - (15 * times_copy)) / times_copy)  # 方框間隔
        if space_Interval > 5:
            rect_width = 15  # 方框寬度
            space_Interval = 5
        elif int(((Length - (bleeding * 2) / layout) - 10 - (10 * times_copy)) / times_copy) > 3:
            rect_width = 10
            space_Interval = 3
        else:
            rect_width = 7
            space_Interval = 2

        one_ptx = int(Length / layout) - 30
        one_width = (Length - bleeding * 2) / layout

        for i in range(int(page_count)):  # 復制PDF頁
            for time in range(times_copy):
                insertpg = i * (times_input)
                position = i * (times_input) + time
                doc.fullcopy_page(insertpg, position)

        y1 = 0
        y2 = y1 + rect_height
        pn = 0
        for pc in range(page_count):  # 插入方框及頁碼
            for time in range(times_input):
                page1 = doc.load_page(pn)
                for w in range(layout):
                    w_width = one_width * w
                    ptx1 = one_ptx * (w + 1)
                    x1 = int(one_inch * bleeding + 5) + time * rect_width + time * space_Interval + w_width
                    x2 = x1 + rect_width
                    # y2 = y1 + rect_height
                    if bleeding_type[0] == "1":
                        rb1 = fitz.Rect(x1, y1, x2, y2)  # 水平
                        page1.insert_text((ptx1, 0), 'Page : ' + str(pn + 1), fontsize=5, rotate=rotate)  # 水平
                    else:
                        rb1 = fitz.Rect(y1, x1, y2, x2)  # 垂直
                        page1.insert_text((0, ptx1), 'Page : ' + str(pn + 1), fontsize=5, rotate=rotate)  # 垂直

                    shape = page1.new_shape()
                    shape.draw_rect(rb1)
                    shape.finish(width=0.3, fill=(1.0, 0.0784313725490196, 0.5764705882352941))
                    shape.insert_textbox(rb1, str(time + 1), color=(1.0, 1.0, 1.0), fontsize=5, align=1, rotate=rotate)
                    # print(rb1,ptx1)
                    shape.commit()
                pn = pn + 1

        # doc.save("./temp4.pdf", garbage=0, deflate=True)
        return doc
    def insertFXpage(lastfilename,r_tab,rx,rx2,ry):
        """
        插入FX每一類檔案標簽PDF檔
        """
        path, filename = os.path.split(lastfilename)
        lastdoc = fitz.open(lastfilename)
        width, height = lastdoc[0].mediabox_size
        # width, height = lastdoc[0].rect.width, lastdoc[0].rect.height
        # r_tab = []
        # r_tab.append(fitz.Rect(0.0, 0.0, 612.0, 396.0))
        # r_tab.append(fitz.Rect(612.0, 0.0, 1224.0, 396.0))
        # r_tab.append(fitz.Rect(0.0, 396.0, 612.0, 792.0))
        # r_tab.append(fitz.Rect(612.0, 396.0, 1224.0, 792.0))
        # r_tab.append(fitz.Rect(0.0, 792.0, 612.0, 1188.0))
        # r_tab.append(fitz.Rect(612.0, 792.0, 1224.0, 1188.0))
        fontsize=int((r_tab[0][2]-r_tab[0][0])/72*6)        #水平像素
        print("水平像素:",fontsize)
        page1 = lastdoc.new_page(width=width, height=height)

        text = ""
        text2 = ''
        fontfile = r'C:\Windows\Fonts\MICROSS.TTF'
        font = fitz.Font(fontfile=fontfile)
        tw = fitz.TextWriter(page1.rect)
        tw2 = fitz.TextWriter(page1.rect)
        if str(filename).upper().find('_JQ_SAMPLE_') >= 0:
            text = str(filename).split('_')[0] + "\nJQ樣板"
        elif str(filename).upper().find('_QC_SAMPLE_') >= 0:
            text = str(filename).split('_')[0] + "\nsignoff記錄"
        elif str(filename).upper().find('_SAMPLE_') >= 0:
            text = str(filename).split('_')[0] + "\n大貨樣板"
        elif str(filename).upper().find('_MAIN_') >= 0:
            text = str(filename).split('_')[0] + "\n大貨換單"
            text2 = str(filename).split('_')[0] + "\n大貨換單"
        else:
            name = filename.split('_')
            l = len(name)
            for s in range(l):
                if s != 1 and s != 2 and s != l - 1 and s != l - 2:
                    text = text + name[s] + " "
            text = text.strip()

        for i in r_tab:
            # tw.append((i[0]+20, i[1]+20), str(i) + text, font=font)  # is ok
            revi = fitz.Rect(i[0], i[1] + (i[3] - i[1]) / 2, i[2], i[3])
            tw.fill_textbox(revi, text, align=fitz.TEXT_ALIGN_CENTER, font=font, fontsize=fontsize)
            tw2.fill_textbox(revi, text2, align=fitz.TEXT_ALIGN_CENTER, font=font, fontsize=fontsize)
        tw.write_text(page1)
        page2 = lastdoc.new_page(width=width, height=height)
        tw.write_text(page2)
        if str(filename).upper().find('_MAIN_') >= 0:
            page3 = lastdoc.new_page(width=width, height=height)
            page3.draw_rect(rx, color=colx, fill=colx, overlay=False)
            page3.draw_rect(rx2, color=colx, fill=colx, overlay=False)
            page3.draw_rect(ry, color=coly, fill=coly, overlay=False)
            tw2.write_text(page3)
            page4 = lastdoc.new_page(width=width, height=height)
            page4.draw_rect(rx, color=colx, fill=colx, overlay=False)
            page4.draw_rect(rx2, color=colx, fill=colx, overlay=False)
            page4.draw_rect(ry, color=coly, fill=coly, overlay=False)
            tw2.write_text(page4)
            page5 = lastdoc.new_page(width=width, height=height)
            page5.draw_rect(rx, color=colx, fill=colx, overlay=False)
            page5.draw_rect(rx2, color=colx, fill=colx, overlay=False)
            page5.draw_rect(ry, color=coly, fill=coly, overlay=False)
            tw2.write_text(page5)
            page6 = lastdoc.new_page(width=width, height=height)
            page6.draw_rect(rx, color=colx, fill=colx, overlay=False)
            page6.draw_rect(rx2, color=colx, fill=colx, overlay=False)
            page6.draw_rect(ry, color=coly, fill=coly, overlay=False)
            tw2.write_text(page6)

            page7 = lastdoc.new_page(width=width, height=height)
            page7.draw_rect(rx, color=colx, fill=colx, overlay=False)
            page7.draw_rect(rx2, color=colx, fill=colx, overlay=False)
            page7.draw_rect(ry, color=coly, fill=coly, overlay=False)
            tw2.write_text(page7)
            page8 = lastdoc.new_page(width=width, height=height)
            page8.draw_rect(rx, color=colx, fill=colx, overlay=False)
            page8.draw_rect(rx2, color=colx, fill=colx, overlay=False)
            page8.draw_rect(ry, color=coly, fill=coly, overlay=False)
            tw2.write_text(page8)

        lastdoc.saveIncr()

    lastfilename=''
    # path, print_type, direction_type, bleeding_type, e1, e2, e3, e4, e5, e6, e7, e8
    path=list1[15]      #文件路徑

    print_type=list1[1]     #編號次序
    direction_type=list1[2]     #PDF輸出方向
    bleeding_type = list1[3]    #出血位及顯示位置
    e1 = list1[4]   # 橫排個數
    e2 = list1[5]   #橫刀位(寸)
    e3 = list1[6]   #左右出血位(寸)
    e4 = list1[8]   #豎排個數
    e5 = list1[9]   #豎刀位(寸)
    e6 = list1[7]   #上下出血位(寸)
    e7 = list1[12]  #一箱張數
    e8 = list1[13]  #一個PDF檔頁數

    e0 = list1[0]   #文件名
    r2c3 = list1[10]  #實際成品寬
    r2c4 = list1[11]  #實際成品高
    e14 = list1[14]  # PDF 总页数
    r3c3 = float(list1[16])  # 水平位置
    r3c4 = float(list1[17])  # 垂直位置
    check_Double_sided = list1[18]  # 是否雙面, 文字: "True", "False"
    r1c5 = float(list1[19])  # 出血位右邊增減
    r1c6 = float(list1[20])  # 出血位下邊增減
    addart = list1[21]   # 是否有art 稿件, 沒有就是no_art
    check_addblank = list1[22]  # 是否加隔紙, 文字: "True", "False"
    check_addblank2 = list1[23]  # 是否單面加空白底頁, 文字: "True", "False"
    PDF_rotate = int(list1[24])  #PDF角度
    addpagenum = list1[25]  # 是否加頁碼
    times_input = int(list1[26])  # 重復數量

    times_copy = times_input - 1  # 實際復制次數
    startime = datetime.datetime.now()
    ###參數
    one_inch = 72  # 1寸=72像素, 常量

    # width_layout = 1  # 橫排個數
    # interval_w = 0.25  # 橫刀位(寸)
    # bleeding_w = 0.5  # 左右出血位(寸)
    # height_layout = 3  # 豎排個數
    # interval_h = 0.25  # 豎刀位(寸)
    # bleeding_H = 0.5  # 上下出血位(寸)
    # cnt_page = 1000  # PDF 排版後的張數
    # pdf_file_page_qty = 3000  # 每個PDF檔輸出的頁數

    width_layout = int(e1)  # 橫排個數
    interval_w = float(e2)  # 橫刀位(寸)
    bleeding_w = float(e3)  # 左右出血位(寸)
    height_layout = int(e4)  # 豎排個數
    interval_h = float(e5)  # 豎刀位(寸)
    bleeding_H = float(e6)  # 上下出血位(寸)
    cnt_page = int(e7)  # PDF 排版後的張數
    pdf_file_page_qty = int(e8)  # 一個PDF檔頁數
    # cnt_qty=int(e8/e7) #一個PDF檔有多少箱


    ###參數

    # if divmod(pdf_file_page_qty, cnt_page)[1] != 0:
    #     print('PDF頁數要整箱數量!')
    #     sys.exit(0)

    Layout = width_layout * height_layout  # 排版個數
    width_interval = one_inch * interval_w  # 寬間隔按像素
    height_interval = one_inch * interval_h  # 高間隔按像素
    width_bleeding = one_inch * bleeding_w  # 左右邊出血位按像素
    height_bleeding = one_inch * bleeding_H  # 上下邊出血位按像素
    catron_qty = cnt_page * Layout  # 一箱記錄數量
    pdf_file_qty = pdf_file_page_qty * Layout  # 一個PDF檔記錄數量

    # scanfold=r'D:\4081933-2 PPF\PDF\OE\IMAGE\L4'
    scanfold,select_filename = os.path.split(path)
    PDF_filename=os.path.splitext(select_filename)[0]
    doc = fitz.open()  # empty output PDF

    t1 = []
    for s in os.listdir(scanfold):
        checkname=re.findall(PDF_filename+r'_C1_\d{7}-\d{7}\.pdf',s,re.S)
        if checkname:
            pass
            QMessageBox.information(ui_mainwindow, "錯誤提示!", "PDF檔: "+s+"已處理! 程式不會再處理此檔案. " + chr(10)+"如需要重新處理, 請先刪除已處理的PDF檔.")
            return
        if s.startswith(PDF_filename) and s==select_filename:
            t2 = (1, s)
            t1.append(t2)
        if s.startswith(PDF_filename) and s!=select_filename:
            num1=int(s[len(PDF_filename) + 4:len(s) - 4])
            t2 = (num1, s)
            t1.append(t2)

    t0 = sorted(t1, key=lambda x: x[0])

    # width=0
    # height=0

    print(len(t0))

    last_infile = scanfold + "\\" + t0[len(t0) - 1][1]
    src = fitz.open(last_infile)
    page_count_last = src.page_count

    first_infile = scanfold + "\\" + t0[0][1]
    src = fitz.open(first_infile)
    page_count_first = src.page_count

    width, height = src[0].mediabox_size
    # width, height = src[0].rect.width, src[0].rect.height
    cropbox=src[0].rect
    print("cropbox:",cropbox)

    cropbox[0] = float(r3c3)*one_inch
    cropbox[1] = float(r3c4)*one_inch
    cropbox[2] = float(r2c3)*one_inch
    cropbox[3] = float(r2c4)*one_inch

    width_cropbox, height_cropbox = src[0].rect.br
    if width_cropbox!=width or height_cropbox!=height:
        width=width_cropbox
        height=height_cropbox

    if cropbox[2]!=width or cropbox[3]!=height:
        width = cropbox[2]
        height = cropbox[3]

    print("width,height: ",width/72,height/72)
    if PDF_rotate == 90 or PDF_rotate == 270:
        width, height = height, width   #對調長寬
    r_tab = []
    x1 = 0
    y1 = 0
    x1 = width_bleeding
    y1 = height_bleeding
    if direction_type[0]=='1':  #PDF輸出方向, 1 是從左到右
        for hl in range(0, height_layout):
            y2 = y1 + height
            for wl in range(0, width_layout):
                x2 = x1 + width
                r_tab.append(fitz.Rect(x1, y1, x2, y2))
                tuple_layout = ()
                x1 = x2+ width_interval
            x1 = width_bleeding
            y1 = y2 + height_interval
        print("r_tab1:",r_tab)
    else:
        for wl in range(0, width_layout):
            x2 = x1 + width
            for hl in range(0, height_layout):
                y2 = y1 + height
                r_tab.append(fitz.Rect(x1, y1, x2, y2))
                tuple_layout = ()
                y1 = y2 + height_interval
            x1 = x2+width_interval
            y1 = height_bleeding

    data_count = page_count_first * (len(t0) - 1) + page_count_last

    # scanfold, select_filename = os.path.split(path)
    # list1 = select_filename.split(" ")
    # list2 = [x for x in list1 if x != '']
    # l2="1" + list2[0][2:].replace("-", "")


    std_pdf_file_count1, std_end_count = divmod(data_count, pdf_file_qty)
    std_pdf_file_count2, end_count = divmod(std_end_count, catron_qty)
    last_cnt_seq = data_count - end_count  # 最後一箱的編號
    # catron_qty= cnt_page * Layout_qty   #一箱記錄數量
    # pdf_file_qty=pdf_file_page_qty*Layout_qty     #一檔PDF檔記錄數量

    if print_type.startswith('1') and check_Double_sided == "False":
        list1,list2_start,list2_end,cnt_page_end = sort_data1(pdf_file_qty, end_count, Layout, cnt_page, std_pdf_file_count2)
    elif print_type.startswith('2') and check_Double_sided == "False":
        list1, list2_start, list2_end, cnt_page_end = sort_data2(pdf_file_qty, end_count, Layout, cnt_page, std_pdf_file_count2)
    elif print_type.startswith('3') and check_Double_sided == "False":
        list1, list2_start, list2_end, cnt_page_end = sort_data3(pdf_file_qty, end_count, Layout, cnt_page, std_pdf_file_count2)
    elif print_type.startswith('1') and check_Double_sided == "True":
        list1, list2_start, list2_end, cnt_page_end = sort_data1_double(std_pdf_file_count1,end_count, width_layout, height_layout, cnt_page, std_pdf_file_count2,pdf_file_page_qty)
    elif print_type.startswith('2') and check_Double_sided == "True":
        list1, list2_start, list2_end, cnt_page_end = sort_data2_double(std_pdf_file_count1, end_count, width_layout, height_layout, cnt_page, std_pdf_file_count2,pdf_file_page_qty)
    elif print_type.startswith('4') and check_Double_sided == "False":
        list1, list2_start, list2_end, cnt_page_end = sort_data1_fx1(width_layout, height_layout, end_count, cnt_page, pdf_file_page_qty,std_pdf_file_count2)
    elif print_type.startswith('5') and check_Double_sided == "False":
        list1, list2_start, list2_end, cnt_page_end = sort_data1_fx2(width_layout, height_layout, end_count, cnt_page, pdf_file_page_qty,std_pdf_file_count2)
    # for l3 in list3:
    #     print(l3)
    # page1=1
    # page2=1
    rec1 = 1
    rec2 = 0

    pdf_file_seq = 0  # 保存PDF檔順序號
    seq_num = 0  # 順序號
    Page_Total = math.ceil(data_count / Layout)
    width_total = (width * width_layout) + (width_interval * (width_layout-1)) + (width_bleeding * 2)+(r1c5*one_inch)
    height_total = (height * height_layout) + (height_interval * (height_layout-1)) + (height_bleeding * 2)+(r1c6*one_inch)
    print("pagesize:",width_total/72,height_total/72)
    # write_lotus_LaserRoom(path, print_type, bleeding_type, e1, e2, e3, e4, e5, e6, e7, e8, data_count, page_count_first, page_count_last,width,height,width_total,height_total,direction_type,r2c3,r2c4,r3c3,r3c4,r1c5,r1c6,check_Double_sided,addart,check_addblank,check_addblank2,PDF_rotate,addpagenum,times_input)  # 更新Lotus
    insert_tex_type = bleeding_type[0]
    if insert_tex_type == '1':  # 左上角水平位置
        point_x = width_bleeding + width_bleeding
        point_y = height_bleeding / 3
        text_rotate = 0
    elif insert_tex_type == '2':  # 左上角垂直位置
        point_x = width_bleeding / 4
        point_y = height_bleeding + height_bleeding
        text_rotate = 270
    elif insert_tex_type == '3':  # 右上角水平位置
        point_x = width_total - 100  # 總寬度100像素
        point_y = height_bleeding / 3
        text_rotate = 0

    elif insert_tex_type == '4':  # 右上角垂直位置
        point_x = width_total - (width_bleeding / 3)
        point_y = height_bleeding + height_bleeding
        text_rotate = 270

    elif insert_tex_type == '5':  # 左下角水平位置
        point_x = width_bleeding + width_bleeding
        point_y = height_total - (height_bleeding / 4)
        text_rotate = 0

    elif insert_tex_type == '6':  # 左下角垂直位置
        point_x = width_bleeding / 3
        point_y = height_total - 100
        text_rotate = 90

    elif insert_tex_type == '7':  # 右下角水平位置
        point_x = width_total - 100  # 總寬度100像素
        point_y = height_total - (height_bleeding / 4)
        text_rotate = 0

    elif insert_tex_type == '8':  # 右下角垂直位置
        point_x = width_total - (width_bleeding / 4)
        point_y = height_total - 100
        text_rotate = 90

    if addart != 'no_art':
        if check_addblank == 'True' and check_addblank2 =='True':
            addblank_page=cnt_page*2*times_input
        else:
            addblank_page = cnt_page

        background = fitz.open(addart)
        widthbg, heightbg = background[0].mediabox_size
        # widthbg, heightbg  = background[0].rect.width, background[0].rect.height
        # print("widthbg,heightbg:", widthbg,heightbg)
        # 72像素=1寸,36像素=0.5寸
        r1 = fitz.Rect(0, 0, widthbg, heightbg)
        rx = fitz.Rect(widthbg / 4, 36, widthbg / 4 + 20, heightbg-36)  # 矩陣-豎向
        rx2 = fitz.Rect(widthbg / 1.3, 36, widthbg / 1.3 + 20, heightbg - 36)  # 矩陣-豎向
        ry = fitz.Rect(36, heightbg / 3, widthbg-36, heightbg / 3 + 20)  # 矩陣-橫向
        colx = fitz.utils.getColor("DodgerBlue")  # 天藍: DodgerBlue
        coly = fitz.utils.getColor("magenta")  # 洋紅: magenta



    for l1 in range(1,pdf_file_page_qty+1):     #标准PDF档数量
        if addart != 'no_art':  #如果有稿件.
            doc.new_page(-1, width=widthbg, height=heightbg)
            if background.page_count > 1 and check_addblank2 =='False':  # 稿件有兩頁并且不是單面加稿件底頁
                if divmod(l1, 2)[1] == 1:
                    doc[l1-1].show_pdf_page(r1, background, 0, keep_proportion=1)
                else:
                    doc[l1-1].show_pdf_page(r1, background, 1, keep_proportion=1)
            else:   # 稿件只有一頁.
                doc[l1-1].show_pdf_page(r1, background, 0, keep_proportion=1)
        else:
            doc.new_page(-1, width=width_total, height=height_total)


        # print(int(point_x), int(point_y), text_rotate)
        if insert_tex_type != '0':
            # page1.insert_text((int(point_x), int(point_y)), 'Page: ' + str(l1), fontsize=7, rotate=text_rotate)
            test=1
        # fontname = 'Arial', fontsize = 7,
        # page.append(page_new)
    page_seq2 = std_pdf_file_count1 * pdf_file_page_qty + 1     #够一箱， 但不够PDF档数量的页码
    page_seq3 = std_pdf_file_count1 * pdf_file_page_qty + std_pdf_file_count2 * cnt_page + 1    ##尾箱数量的页码
    if std_pdf_file_count2 > 0:     #够一箱， 但不够PDF档数量
        doc2 = fitz.open()  # empty output PDF
        # print("够一箱， 但不够PDF档数量:",(cnt_page*std_pdf_file_count2)+1)
        for l2 in range(1,(cnt_page*std_pdf_file_count2)+1):
            # page2 = doc2.new_page(-1, width=width_total, height=height_total)
            if addart != 'no_art':  # 如果有稿件.
                page2=doc2.new_page(-1, width=widthbg, height=heightbg)
                if background.page_count > 1 and check_addblank2 =='False':  # 稿件有兩頁并且不是單面加稿件底頁
                    if divmod(l2, 2)[1] == 1:
                        doc2[l2 - 1].show_pdf_page(r1, background, 0, keep_proportion=1)
                    else:
                        doc2[l2 - 1].show_pdf_page(r1, background, 1, keep_proportion=1)
                else:  # 稿件只有一頁.
                    doc2[l2 - 1].show_pdf_page(r1, background, 0, keep_proportion=1)
            else:
                page2=doc2.new_page(-1, width=width_total, height=height_total)

            # if insert_tex_type != '0':
            if addpagenum == 'True':
                page2.insert_text((int(point_x), int(point_y)), 'Page: ' + str(page_seq2), fontsize=6, rotate=text_rotate)
                page_seq2=page_seq2+1
        doc2.save("./temp2.pdf", garbage=0, deflate=True)

    if end_count > 0:   #尾箱数量
        doc3 = fitz.open()  # empty output PDF
        for l3 in range(1,cnt_page_end+1):
            # page3 = doc3.new_page(-1, width=width_total, height=height_total)
            if addart != 'no_art':  # 如果有稿件.
                page3=doc3.new_page(-1, width=widthbg, height=heightbg)
                if background.page_count > 1 and check_addblank2 =='False':  # 稿件有兩頁并且不是單面加稿件底頁
                    if divmod(l3, 2)[1] == 1:
                        doc3[l3 - 1].show_pdf_page(r1, background, 0, keep_proportion=1)
                    else:
                        doc3[l3 - 1].show_pdf_page(r1, background, 1, keep_proportion=1)
                else:  # 稿件只有一頁.
                    doc3[l3 - 1].show_pdf_page(r1, background, 0, keep_proportion=1)
            else:
                page3 = doc3.new_page(-1, width=width_total, height=height_total)

            # if insert_tex_type != '0':
            if addpagenum == 'True':
                page3.insert_text((int(point_x), int(point_y)), 'Page: ' + str(page_seq3), fontsize=6, rotate=text_rotate)
                page_seq3=page_seq3+1
        doc3.save("./temp3.pdf", garbage=0, deflate=True)

    doc.save("./temp.pdf", garbage=0, deflate=True)





    # target = './' + PDF_filename + "_0000001-" + str(pdf_file_qty).rjust(7, '0') + ".pdf"
    # target = outputfile + PDF_filename +"/" + PDF_filename + "_0000001-" + str(pdf_file_qty).rjust(7, '0') + ".pdf"
    # if check_addblank == 'True':
    #40 PDF文件改為10
    # fx=str(PDF_filename).split(" ")
    # Hjob = "1" + fx[0].replace("-", "")[2:8]
    # Hlot = fx[1][1:]
    # Hitem = fx[2]
    # fx3=''
    # for i in fx[3:]:
    #     fx3 = fx3 + "_" + i
    # PDF_filename_output = Hjob + Hitem + "P" + Hlot + "_L" + Hlot + "_" + Hitem + fx3

    # if ui_mainwindow.Processing_department == '打印部':
    #     scanfold_W=scanfold+"_W"
    #     process_dir = Path(scanfold_W)
    #     if not Path(scanfold_W).is_dir():
    #         os.makedirs(process_dir)
    # else:
    #     scanfold_W=scanfold
    # 40 PDF文件改為10
    # 這個是補數用的.
    PDF_filename_output = PDF_filename
    scanfold_W=ui_mainwindow.outputfile
    # 這個是補數用的.

    target = scanfold_W +"/" + PDF_filename_output + "_C1_0000001-" + str(pdf_file_qty).rjust(7, '0') + ".pdf"

    if data_count>=pdf_file_qty:
        copyfile('./temp.pdf', target)
        doc_process = fitz.open(target)
    elif data_count>=catron_qty:
        # doc_process=doc2
        target = scanfold_W +"/" + PDF_filename_output + "_"+"C" + str(pdf_file_seq+1)+"_" + str(1).rjust(7, '0') + '-' + str(cnt_page*Layout*std_pdf_file_count2).rjust(7, '0') + ".pdf"
        dele_PDF = scanfold_W +"/" + PDF_filename_output + "_"+"C" + str(pdf_file_seq+1)+"_" + str(1).rjust(7, '0') + '-' + str(cnt_page*Layout*std_pdf_file_count2).rjust(7, '0') + ".pdf"
        lastfilename = scanfold_W +"/" + PDF_filename_output + "_"+"C" + str(pdf_file_seq+1)+"_" + str(1).rjust(7, '0') + '-' + str(cnt_page*Layout*std_pdf_file_count2).rjust(7, '0') + ".pdf"
        copyfile('./temp2.pdf', target)
        doc_process = fitz.open(target)
        list1 = list2_start
    else:
        # doc_process = doc3
        target = scanfold_W + "/" + PDF_filename_output + "_" + "C" + str(pdf_file_seq + 1) + "_" + str(1).rjust(7,'0') + '-' + str(data_count).rjust(7, '0') + ".pdf"
        lastfilename = scanfold_W + "/" + PDF_filename_output + "_" + "C" + str(pdf_file_seq + 1) + "_" + str(1).rjust(7, '0') + '-' + str(data_count).rjust(7, '0') + ".pdf"
        copyfile('./temp3.pdf', target)
        doc_process = fitz.open(target)
        list1 = list2_end
    page_seq=1
    for k1 in t0:


        # t=k.lower().find("main")
        # print(t)
        # if k.lower().find("main")>-1 and k.lower().find("job")==-1:     #掃描第一個PDF檔, 取總頁數
        infile = scanfold + "\\" + k1[1]
        src = fitz.open(infile)

        for spage in src:
            # insert input page into the correct rectangle
            # if rec2>last_cnt_seq:
            #     l3 = list2_end[seq_num]
            # else:
            #     l3 = list1[seq_num]
            # print("check:",seq_num)
            l3 = list1[seq_num]
            page_number = l3[3]

            layout_number = l3[4]
            # print('page_number:', page_number, "layout_number:", layout_number )
            # print("Check page_number: ",page_number - 1, "Check spage.number: ", spage.number)
            doc_process.load_page(page_number - 1).show_pdf_page(r_tab[layout_number - 1], src, spage.number, keep_proportion=1, clip=cropbox, rotate=PDF_rotate)  # 1 select output rect, #2 input document, #3 input page number
            rec2 = rec2 + 1
            seq_num = seq_num + 1
            # check_save=rec2 % pdf_file_qty
            if rec2 % pdf_file_qty == 0:    #標準箱整個檔案處理

                print('最後輸出記錄號 std: ', rec1, rec2, pdf_file_seq)

                if times_input>1:     #重復多少張
                    doc_process=doutimes(doc_process)
                if addpagenum == 'True':    #加頁碼
                # if insert_tex_type != '0':
                    for pn in doc_process:
                        pn.insert_text((int(point_x), int(point_y)), 'Page_seq: ' + str(page_seq), fontsize=7, rotate=text_rotate)
                        page_seq=page_seq+1
                if check_addblank2 =='True':    #單面加空白底頁
                    pagecount=doc_process.page_count
                    for i in range(int(pagecount)):
                        insertpg = i * 2 + 1
                        # print("insertpg:", insertpg)
                        doc_process.insert_page(insertpg, width=width_total, height=height_total)
                        page1 = doc_process.load_page(insertpg)
                        page1.show_pdf_page(r1, background, 1, keep_proportion=1)
                if check_addblank == 'True':    #加隔紙
                    for i in range(int(doc_process.page_count / addblank_page)):
                        insertpg = int(addblank_page * i) + i*2 + addblank_page
                        print("insertpg:", insertpg)
                        doc_process.insert_page(insertpg, text='blank1', width=widthbg, height=heightbg)
                        pageblank1 = doc_process.load_page(insertpg)
                        pageblank1.draw_rect(rx, color=colx, fill=colx, overlay=False)
                        pageblank1.draw_rect(rx2, color=colx, fill=colx, overlay=False)
                        pageblank1.draw_rect(ry, color=coly, fill=coly, overlay=False)

                        doc_process.insert_page(insertpg+1, text='blank2', width=widthbg, height=heightbg)
                        pageblank2 = doc_process.load_page(insertpg + 1)
                        pageblank2.draw_rect(rx, color=colx, fill=colx, overlay=False)
                        pageblank2.draw_rect(rx2, color=colx, fill=colx, overlay=False)
                        pageblank2.draw_rect(ry, color=coly, fill=coly, overlay=False)
                doc_process.saveIncr()
                doc_process.close()
                pdf_file_seq = pdf_file_seq + 1
                rec1 = rec2 + 1
                seq_num = 0
                if pdf_file_seq != std_pdf_file_count1:
                    # target = './' + PDF_filename + "_" + str(rec1).rjust(7, '0') + '-' + str(rec2 + pdf_file_qty).rjust(7, '0') + ".pdf"
                    # target = outputfile + PDF_filename +"/" + PDF_filename + "_" + str(rec1).rjust(7, '0')+'-'+str(rec2+pdf_file_qty).rjust(7, '0') + ".pdf"

                    target = scanfold_W +"/" + PDF_filename_output + "_"+"C" + str(pdf_file_seq+1)+"_" + str(rec1).rjust(7, '0') + '-' + str(rec2 + pdf_file_qty).rjust(7, '0') + ".pdf"
                    lastfilename = scanfold_W +"/" + PDF_filename_output + "_"+"C" + str(pdf_file_seq+1)+"_" + str(rec1).rjust(7, '0') + '-' + str(rec2 + pdf_file_qty).rjust(7, '0') + ".pdf"
                    print("lastfilename:",lastfilename)
                    copyfile('./temp.pdf', target)
                    doc_process = fitz.open(target)
                elif std_pdf_file_count2>0:
                    # doc_process = doc2
                    target = scanfold_W +"/" + PDF_filename_output + "_"+"C" + str(pdf_file_seq+1)+"_" + str(rec1).rjust(7, '0') + '-' + str(rec2 + cnt_page*Layout*std_pdf_file_count2).rjust(7, '0') + ".pdf"
                    dele_PDF = scanfold_W +"/" + PDF_filename_output + "_"+"C" + str(pdf_file_seq+1)+"_" + str(rec1).rjust(7, '0') + '-' + str(rec2 + cnt_page*Layout*std_pdf_file_count2).rjust(7, '0') + ".pdf"
                    lastfilename = scanfold_W +"/" + PDF_filename_output + "_"+"C" + str(pdf_file_seq+1)+"_" + str(rec1).rjust(7, '0') + '-' + str(rec2 + cnt_page*Layout*std_pdf_file_count2).rjust(7, '0') + ".pdf"
                    copyfile('./temp2.pdf', target)
                    doc_process = fitz.open(target)
                elif end_count==0:
                    pass #總數剛好整箱.
                else:

                    # doc_process = doc3
                    target = scanfold_W +"/" + PDF_filename_output + "_"+"C" + str(pdf_file_seq+1)+"_" + str(rec1).rjust(7, '0') + '-' + str(rec2 + end_count).rjust(7, '0') + ".pdf"
                    lastfilename = scanfold_W +"/" + PDF_filename_output + "_"+"C" + str(pdf_file_seq+1)+"_" + str(rec1).rjust(7, '0') + '-' + str(rec2 + end_count).rjust(7, '0') + ".pdf"
                    copyfile('./temp3.pdf', target)
                    doc_process = fitz.open(target)
                    list1 = list2_end

            if std_pdf_file_count2 > 0 and rec2 == last_cnt_seq:    #標準箱但不夠一整個檔案

                print('最後輸出記錄號 cnt end: ', rec1, rec2, pdf_file_seq)
                # doc_process.save('./'+PDF_filename+"_" + str(pdf_file_seq).rjust(4, '0') + '.pdf', garbage=3, deflate=True)
                # doc_process.save('./' + PDF_filename + "_" + str(rec1).rjust(7, '0')+'-'+str(rec2).rjust(7, '0') + '.pdf', garbage=3, deflate=True)
                # doc_process.save(outputfile + PDF_filename +"/"+ PDF_filename + "_" + str(rec1).rjust(7, '0') + '-' + str(rec2).rjust(7, '0') + '.pdf', garbage=3, deflate=True)
                if times_input>1:     #重復多少張
                    doc_process=doutimes(doc_process)
                if check_addblank2 =='True':    #單面加空白底頁
                    pagecount=doc_process.page_count
                    for i in range(int(pagecount)):
                        insertpg = i * 2 + 1
                        # print("insertpg:", insertpg)
                        doc_process.insert_page(insertpg, width=width_total, height=height_total)
                        page1 = doc_process.load_page(insertpg)
                        page1.show_pdf_page(r1, background, 1, keep_proportion=1)
                if check_addblank == 'True':    #加隔紙
                    for i in range(int(doc_process.page_count / addblank_page)):
                        insertpg = int(addblank_page * i) + i*2 + addblank_page
                        print("insertpg:", insertpg)
                        doc_process.insert_page(insertpg, text='blank1', width=widthbg, height=heightbg)
                        pageblank1 = doc_process.load_page(insertpg)
                        pageblank1.draw_rect(rx, color=colx, fill=colx, overlay=False)
                        pageblank1.draw_rect(rx2, color=colx, fill=colx, overlay=False)
                        pageblank1.draw_rect(ry, color=coly, fill=coly, overlay=False)

                        doc_process.insert_page(insertpg+1, text='blank2', width=widthbg, height=heightbg)
                        pageblank2 = doc_process.load_page(insertpg + 1)
                        pageblank2.draw_rect(rx, color=colx, fill=colx, overlay=False)
                        pageblank2.draw_rect(rx2, color=colx, fill=colx, overlay=False)
                        pageblank2.draw_rect(ry, color=coly, fill=coly, overlay=False)

                doc_process.saveIncr()

                # doc_process.save(scanfold_W + "/" + PDF_filename_output + "_" + "C" + str(pdf_file_seq+1)+"_" + str(rec1).rjust(7, '0') + '-' + str(rec2).rjust(7, '0') + '.pdf', garbage=3, deflate=True)

                # dele_PDF=scanfold_W + "/" + PDF_filename_output + "_" + "C" + str(pdf_file_seq+1)+"_" + str(rec1).rjust(7, '0') + '-' + str(rec2).rjust(7, '0') + '.pdf'
                lastfilename=scanfold_W + "/" + PDF_filename_output + "_" + "C" + str(pdf_file_seq + 1) + "_" + str(rec1).rjust(7, '0') + '-' + str(rec2).rjust(7, '0') + '.pdf'
                print("lastfilename2:", lastfilename)
                rec1_w=deepcopy(rec1)
                print("outputfile1:", outputfile + PDF_filename +"/"+ PDF_filename + "_" + str(rec1).rjust(7, '0') + '-' + str(rec2).rjust(7, '0') + '.pdf' )
                # doc_process.saveIncr()
                doc_process.close()
                pdf_file_seq = pdf_file_seq + 1
                rec1 = rec2 + 1
                seq_num = 0
                # if len(list2_end)>0:    #如果有尾箱.
                if end_count>0:          # 如果有尾箱.
                    # doc_process = doc3
                    target = scanfold_W +"/" + PDF_filename_output + "_"+"C" + str(pdf_file_seq+1)+"_" + str(rec1).rjust(7, '0') + '-' + str(rec2 + end_count).rjust(7, '0') + ".pdf"
                    lastfilename = scanfold_W +"/" + PDF_filename_output + "_"+"C" + str(pdf_file_seq+1)+"_" + str(rec1).rjust(7, '0') + '-' + str(rec2 + end_count).rjust(7, '0') + ".pdf"
                    copyfile('./temp3.pdf', target)
                    doc_process = fitz.open(target)
                    list1 = list2_end  # 最後尾箱列表

            # if rec2 == data_count and len(list2_end) > 0:
            if rec2 == data_count and end_count>0:   # 尾箱.

                print('最後輸出記錄號 end: ', rec1, rec2,pdf_file_seq)
                # doc_process.saveIncr()
                # doc_process.save('./'+PDF_filename+"_" + str(pdf_file_seq).rjust(4, '0') + '.pdf', garbage=3, deflate=True)
                # doc_process.save('./' + PDF_filename + "_" + str(rec1).rjust(7, '0')+'-'+str(rec2).rjust(7, '0') + '.pdf', garbage=3, deflate=True)
                # doc_process.save(outputfile + PDF_filename +"/"+PDF_filename + "_" + str(rec1).rjust(7, '0') + '-' + str(rec2).rjust(7, '0') + '.pdf', garbage=3, deflate=True)
                if times_input>1:     #重復多少張
                    doc_process=doutimes(doc_process)
                if check_addblank2 =='True':    #單面加空白底頁
                    pagecount=doc_process.page_count
                    for i in range(int(pagecount)):
                        insertpg = i * 2 + 1
                        # print("insertpg:", insertpg)
                        doc_process.insert_page(insertpg, width=width_total, height=height_total)
                        page1 = doc_process.load_page(insertpg)
                        page1.show_pdf_page(r1, background, 1, keep_proportion=1)

                doc_process.saveIncr()
                doc_process.close()
                lastfilename = scanfold_W + "/" + PDF_filename_output + "_" + "C" + str(pdf_file_seq + 1) + "_" + str(rec1).rjust(7, '0') + '-' + str(rec2).rjust(7, '0') + '.pdf'
                # 不做合並, 因為處理時間太慢.
                # if std_pdf_file_count2>0:        #標準箱但不夠一整個檔案
                #     print('dele_PDF:', dele_PDF)
                #     doc_process2 = fitz.open(dele_PDF)
                #     doc_process = fitz.open(lastfilename)
                #     doc_process2.insert_pdf(doc_process)
                #     doc_process2.saveIncr()
                #     doc_process2.close()
                #     doc_process.close()
                #
                #     os.remove(lastfilename)
                #     rec1=rec1_w
                #     pdf_file_seq=pdf_file_seq-1

                # doc_process.saveIncr()
                # doc_process.save(scanfold_W + "/" + PDF_filename_output + "_" + "C" + str(pdf_file_seq + 1) + "_" + str(rec1).rjust(7, '0') + '-' + str(rec2).rjust(7, '0') + '.pdf', garbage=3, deflate=True)
                # lastfilename=scanfold_W + "/" + PDF_filename_output + "_" + "C" + str(pdf_file_seq + 1) + "_" + str(rec1).rjust(7, '0') + '-' + str(rec2).rjust(7, '0') + '.pdf'
                print("lastfilename3:", lastfilename)
                print("outputfile3:",outputfile + PDF_filename +"/"+ PDF_filename + "_" + str(rec1).rjust(7, '0') + '-' + str(rec2).rjust(7, '0') + '.pdf')
                # QMessageBox.information(ui_mainwindow,"消息",outputfile + PDF_filename + "_" + str(rec1).rjust(7, '0') + '-' + str(rec2).rjust(7, '0') + '.pdf')

                # doc_process.close()
                # if std_pdf_file_count2 > 0:
                #     print('rename:',target,lastfilename)
                    # os.rename(target,lastfilename)
                pdf_file_seq = 0

        src.close()
    # if check_addblank == 'True':        #加隔紙,即FX類型, 需要在最后的檔案插入標簽
    #     print("lastfilename4:", lastfilename)
    #     insertFXpage(lastfilename,r_tab,rx,rx2,ry)



    end_time = datetime.datetime.now()

    # show_infor.set(" 排版處理完成!")
    with open("./PDF_logfile.txt", 'a+', encoding='UTF-8') as log:
        log.write("處理的PDF: "+e0+chr(13))
        log.write("處理時間: " + startime.__str__()+" 至 "+end_time.__str__()+" ; 共用時間: "+(end_time - startime).__str__()+chr(13))
    # print('時間: ', startime, end_time, end_time - startime)


if __name__ == '__main__':
    app = QApplication([])

    outputfile=""
    dirpath = "//172.16.3.6/Laser/Production/PDF"
    model_view2 = QStandardItemModel()
    model_view2.setHorizontalHeaderLabels(
        ["文件名", "編號次序", "PDF輸出方向", "出血位及顯示位置", "橫排個數", "橫刀位", "左右出血位", "上下出血位",
         "豎排個數", "豎刀位", "實際成品寬", "實際成品高", "一箱張數", "PDF檔頁數", "总页数", "檔案路徑","水平位置", "垂直位置","是否雙面","右邊增減","下邊增減","排版PDF稿件","加隔紙","單面加稿件底頁","PDF角度","加頁碼","重復數量"])

    ui_mainwindow = mainwindow(QMainWindow())

    s = win32com.client.Dispatch('Notes.NotesSession')
    db_LM = s.GetDatabase(ui_mainwindow.Lotus_server, "PublicNSF\\LaserMat.nsf")
    view_joblotitem = db_LM.GetView("searchPDFbyjoblotitem")
    one_inch = 72  # 1寸=72像素, 常量
    listwindow = []  # 定義窗口全局變量
    app.exec_()







